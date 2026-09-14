import subprocess
import os
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Alignment
from fpdf import FPDF
import tempfile
import plotly.express as px

# ==========================================
# مکانیزم همگام‌سازی خودکار دیتابیس‌های اکسل با گیت‌هاب
# ==========================================
def auto_sync_github():
    try:
        status = subprocess.run(["git", "status", "--porcelain"], capture_output=True, text=True)
        if status.stdout.strip():
            subprocess.run(["git", "config", "--global", "user.email", "dibaposh@factory.com"], check=True)
            subprocess.run(["git", "config", "--global", "user.name", "DibaPosh System"], check=True)
            subprocess.run(["git", "add", "*.xlsx"], check=True)
            subprocess.run(["git", "commit", "-m", "Auto-sync database files"], check=True)
            subprocess.run(["git", "push", "origin", "main"], check=True)
    except Exception as e:
        pass

# اجرای خودکار در شروع برنامه
auto_sync_github()

st.set_page_config(page_title="دیپاپوش اسپادانا | سامانه جامع فرماندهی صنعتی", page_icon="🧵", layout="wide")

ORDERS_FILE = 'Orders_DibaPosh.xlsx'
INVENTORY_FILE = 'Inventory_DibaPosh.xlsx'
PRODUCTS_MASTER_FILE = 'Products_Master.xlsx'
MPS_FILE = 'Production_Schedule_MPS.xlsx'
QC_FILE = 'Quality_Control_QC.xlsx'
AUDIT_LOG_FILE = 'Audit_Activity_Log.xlsx'
WAREHOUSE_LOG_FILE = 'Warehouse_Transactions_Log.xlsx'
USERS_DB_FILE = 'Users_Database.xlsx'
REAL_PRODUCTS_EXCEL = '1405-06-01-16-16---08.001.xls.xlsx'
METERS_PER_ROLL = 500.0

def safe_init_and_recovery():
    try:
        if not os.path.exists(USERS_DB_FILE):
            default_users = pd.DataFrame([
                {"username": "admin", "password": "123", "role": "مدیرعامل", "name": "مدیرعامل محترم", "permissions": "ثبت سفارش,کارتابل مشتریان,مدیریت سفارشات,برنامه‌ریزی تولید,تولید و کارگاه,انباردار,کنترل کیفیت,گزارشات کلیدی,مدیریت کاربران"},
                {"username": "sales", "password": "123", "role": "مدیر فروش", "name": "آقای اسماعیلی (مسئول فروش)", "permissions": "ثبت سفارش,کارتابل مشتریان,مدیریت سفارشات"}
            ])
            default_users.to_excel(USERS_DB_FILE, index=False)

        if not os.path.exists(PRODUCTS_MASTER_FILE):
            if os.path.exists(REAL_PRODUCTS_EXCEL):
                df_real = pd.read_excel(REAL_PRODUCTS_EXCEL)
                master_data = {
                    'کد کالا': df_real['کد'].values if 'کد' in df_real.columns else ["2601/876", "1201/11", "1201/99"],
                    'محصول / مدل فویل': df_real['نام '].astype(str).str.strip() if 'نام ' in df_real.columns else df_real.iloc[:, 1].astype(str),
                }
                pd.DataFrame(master_data).to_excel(PRODUCTS_MASTER_FILE, index=False)
            else:
                pd.DataFrame(columns=['کد کالا', 'محصول / مدل فویل']).to_excel(PRODUCTS_MASTER_FILE, index=False)

        if not os.path.exists(INVENTORY_FILE):
            if os.path.exists(REAL_PRODUCTS_EXCEL):
                df_real = pd.read_excel(REAL_PRODUCTS_EXCEL)
                inv_data = {
                    'کد کالا': df_real['کد'].values if 'کد' in df_real.columns else ["2601/876", "1201/11", "1201/99"],
                    'محصول / مدل فویل': df_real['نام '].astype(str).str.strip() if 'نام ' in df_real.columns else df_real.iloc[:, 1].astype(str),
                    'موجودی انبار (طاقه)': 0,
                    'موجودی انبار (متر)': 0.0,
                    'رول مادر قواره‌ای (متر)': 0.0,
                    'سایز برش عرض (cm)': 16.0,
                    'طاقه رزرو انبار': 0,
                    'تاریخ انبار خوردن': jdatetime.datetime.now().strftime("%Y/%m/%d"),
                    'مکان قفسه': [f"Rack-{(i%5)+1}-Shelf{(i%3)+1}" for i in range(len(df_real))],
                }
                pd.DataFrame(inv_data).to_excel(INVENTORY_FILE, index=False)
            else:
                pd.DataFrame(columns=['کد کالا', 'محصول / مدل فویل', 'موجودی انبار (طاقه)', 'موجودی انبار (متر)', 'رول مادر قواره‌ای (متر)', 'سایز برش عرض (cm)', 'طاقه رزرو انبار', 'تاریخ انبار خوردن', 'مکان قفسه']).to_excel(INVENTORY_FILE, index=False)
        else:
            df_inv = pd.read_excel(INVENTORY_FILE)
            changed = False
            for col_name, default_val in [('رول مادر قواره‌ای (متر)', 0.0), ('طاقه رزرو انبار', 0), ('سایز برش عرض (cm)', 16.0), ('موجودی انبار (طاقه)', 0), ('موجودی انبار (متر)', 0.0)]:
                if col_name not in df_inv.columns:
                    df_inv[col_name] = default_val
                    changed = True
            if changed:
                df_inv.to_excel(INVENTORY_FILE, index=False)

        if not os.path.exists(ORDERS_FILE):
            pd.DataFrame(columns=[
                'وضعیت', 'شناسه سطر', 'شماره سفارش', 'نام مشتری', 'کد کالا', 'محصول / مدل فویل', 
                'تعداد طاقه', 'متراژ کل (متر)', 'سهم انبار (طاقه)', 'سهم تولید (طاقه)', 'نحوه تامین و ارسال', 'اولویت', 'ثبت‌کننده', 'تاریخ ثبت', 
                'متراژ رول مادر تولیدی', 'تاریخ تولید', 'تاریخ تکمیل تولید', 'تاریخ ارسال', 'تاریخ مرجوعی', 'علت مرجوعی', 'کسری برش (طاقه)'
            ]).to_excel(ORDERS_FILE, index=False)
        else:
            df_ord_chk = pd.read_excel(ORDERS_FILE)
            if 'کسری برش (طاقه)' not in df_ord_chk.columns:
                df_ord_chk['کسری برش (طاقه)'] = 0
                df_ord_chk.to_excel(ORDERS_FILE, index=False)
            if 'کد کالا' not in df_ord_chk.columns:
                df_ord_chk['کد کالا'] = '-'
                df_ord_chk.to_excel(ORDERS_FILE, index=False)

        if not os.path.exists(MPS_FILE):
            pd.DataFrame(columns=[
                'شناسه سطر', 'تاریخ برنامه‌ریزی تولید', 'روز هفته', 'شماره سفارش', 'نام مشتری', 
                'کد کالا', 'محصول / مدل فویل', 'دستگاه تولید', 'طیف رنگی', 'تعداد طاقه تولید', 'متراژ تولید (متر)', 'اولویت', 'تاریخ ثبت سفارش', 'وضعیت اجرای خط', 'یادداشت برنامه‌ریز'
            ]).to_excel(MPS_FILE, index=False)
        else:
            df_mps_chk = pd.read_excel(MPS_FILE)
            if 'کد کالا' not in df_mps_chk.columns:
                df_mps_chk['کد کالا'] = '-'
                df_mps_chk.to_excel(MPS_FILE, index=False)

        if not os.path.exists(QC_FILE):
            pd.DataFrame(columns=[
                'شناسه سطر', 'شماره سفارش', 'نام مشتری', 'کد کالا', 'محصول / مدل فویل', 'تعداد طاقه', 'تاریخ تولید',
                'کیفیت چسب و چسبندگی', 'کیفیت هات‌استمپ شدن', 'مقاومت در برابر سایش', 'ارزیابی ظاهری',
                'نتیجه ارزیابی QC', 'تاریخ ارزیابی', 'مسئول کنترل کیفی', 'توضیحات بازرس'
            ]).to_excel(QC_FILE, index=False)
        else:
            df_qc_chk = pd.read_excel(QC_FILE)
            if 'کد کالا' not in df_qc_chk.columns:
                df_qc_chk['کد کالا'] = '-'
                df_qc_chk.to_excel(QC_FILE, index=False)

        if not os.path.exists(AUDIT_LOG_FILE):
            pd.DataFrame(columns=['تاریخ و ساعت', 'نوع عملیات', 'شرح جزئیات']).to_excel(AUDIT_LOG_FILE, index=False)

        if not os.path.exists(WAREHOUSE_LOG_FILE):
            pd.DataFrame(columns=['تاریخ و ساعت', 'نوع تراکنش', 'کد کالا', 'محصول', 'تعداد طاقه', 'جزئیات']).to_excel(WAREHOUSE_LOG_FILE, index=False)

    except PermissionError:
        st.error("🚨 **خطای دسترسی:** فایل اکسل انبار یا سفارشات باز است. لطفاً آن را ببندید و صفحه را رفرش کنید.")
        st.stop()
    except Exception:
        pass

safe_init_and_recovery()

if 'dark_mode' not in st.session_state:
    st.session_state.dark_mode = False

if st.session_state.dark_mode:
    app_bg = "#0b0f19"
    card_bg = "#131c2e"
    text_color = "#f1f5f9"
    subtext_color = "#94a3b8"
    border_color = "#1e293b"
    input_bg = "#1e293b"
    input_text = "#ffffff"
    tab_list_bg = "#131c2e"
else:
    app_bg = "#f8fafc"
    card_bg = "#ffffff"
    text_color = "#1e293b"
    subtext_color = "#64748b"
    border_color = "#e2e8f0"
    input_bg = "#ffffff"
    input_text = "#0f172a"
    tab_list_bg = "#e2e8f0"

st.markdown(f"""
    <style>
    @import url('https://cdn.jsdelivr.net/gh/fontiran/b-nazanin-font@master/B-Nazanin.css');
    
    .stApp, [data-testid="stAppViewContainer"], [data-testid="stHeader"] {{ 
        background-color: {app_bg} !important; 
        color: {text_color} !important;
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
        font-size: 19px !important;
        font-weight: 700 !important;
    }}
    
    p, span, label, h1, h2, h3, h4, h5, h6, .stMarkdown, .stTextInput, .stSelectbox {{
        color: {text_color} !important;
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
        font-weight: 800 !important;
    }}

    .stSelectbox div[data-baseweb="select"] > div {{
        min-height: 52px !important;
        align-items: center !important;
    }}

    [data-testid="stExpander"] {{
        background-color: {card_bg} !important;
        border: 1px solid {border_color} !important;
        border-radius: 14px !important;
        margin-bottom: 12px;
    }}
    [data-testid="stExpander"] summary {{
        direction: rtl !important;
        text-align: right !important;
        font-size: 19px !important;
        font-weight: 900 !important;
        padding: 12px !important;
    }}
    [data-testid="stExpander"] summary span {{
        display: inline-block;
        width: 100%;
    }}
    [data-testid="stExpander"] svg {{
        position: static !important;
        margin-left: 12px !important;
        order: 2;
    }}
    
    .metric-card {{
        background-color: {card_bg} !important;
        padding: 24px;
        border-radius: 16px;
        border: 2px solid {border_color};
        text-align: center;
        margin-bottom: 15px;
        box-shadow: 0 6px 15px rgba(0, 0, 0, 0.06);
    }}
    
    .landing-header {{
        background: linear-gradient(135deg, #1e3a8a 0%, #2563eb 50%, #3b82f6 100%);
        padding: 40px;
        border-radius: 24px;
        text-align: center;
        box-shadow: 0 15px 30px -5px rgba(37, 99, 235, 0.3);
        border-bottom: 6px solid #eab308;
        margin-bottom: 30px;
        color: white !important;
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
    }}
    .landing-header h1, .landing-header p {{ color: white !important; font-family: 'B Nazanin', Tahoma, sans-serif !important; font-weight: 900 !important; }}
    
    .hero-container {{
        background: linear-gradient(135deg, #1e3a8a 0%, #2563eb 100%);
        padding: 35px;
        border-radius: 20px;
        text-align: center;
        box-shadow: 0 10px 25px -5px rgba(30, 58, 138, 0.25);
        border-bottom: 5px solid #d97706;
        margin-bottom: 25px;
        color: white !important;
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
    }}
    .hero-container h2, .hero-container p {{ color: white !important; font-family: 'B Nazanin', Tahoma, sans-serif !important; font-weight: 800 !important; }}
    
    .company-intro-card, .login-form-card {{
        background-color: {card_bg} !important;
        padding: 40px;
        border-radius: 22px;
        border: 1px solid {border_color};
        box-shadow: 0 12px 30px -10px rgba(0,0,0,0.08);
        height: 100%;
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
    }}
    
    .portal-badge {{
        background: linear-gradient(135deg, rgba(37, 99, 235, 0.1) 0%, rgba(234, 179, 8, 0.1) 100%);
        padding: 14px 18px;
        border-radius: 14px;
        margin-bottom: 14px;
        border: 1px solid {border_color};
        font-weight: 800;
        font-size: 17px;
        display: flex;
        align-items: center;
        gap: 12px;
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
    }}
    
    .notification-banner {{
        background: linear-gradient(135deg, #dc2626 0%, #991b1b 100%);
        color: white;
        padding: 16px 20px;
        border-radius: 14px;
        font-weight: 900;
        font-size: 18px;
        margin-bottom: 20px;
        display: flex;
        align-items: center;
        justify-content: space-between;
        box-shadow: 0 8px 20px rgba(220, 38, 38, 0.3);
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
    }}
    
    .warehouse-alert {{
        background: linear-gradient(135deg, #10b981 0%, #059669 100%);
        color: white;
        padding: 16px 20px;
        border-radius: 14px;
        font-weight: 900;
        font-size: 18px;
        margin-bottom: 20px;
        display: flex;
        align-items: center;
        justify-content: space-between;
        box-shadow: 0 8px 20px rgba(16, 185, 129, 0.3);
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
    }}
    
    .stTabs [data-baseweb="tab-list"] {{ gap: 8px; background-color: {tab_list_bg}; padding: 8px; border-radius: 14px; }}
    .stTabs [data-baseweb="tab"] {{ height: 50px; background-color: {card_bg}; border-radius: 8px; color: {text_color}; font-weight: 900; font-size: 17px; border: 1px solid {border_color}; font-family: 'B Nazanin', Tahoma, sans-serif !important; }}
    .stTabs [aria-selected="true"] {{ background: linear-gradient(135deg, #1e3a8a, #2563eb) !important; color: white !important; border: none !important; }}
    
    input, textarea {{
        background-color: {input_bg} !important;
        color: {input_text} !important;
        border-color: {border_color} !important;
        border-radius: 10px !important;
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
        font-weight: 800 !important;
        font-size: 17px !important;
    }}
    
    .stButton button {{
        background: linear-gradient(135deg, #1e3a8a 0%, #2563eb 100%);
        color: white !important; font-weight: 900; font-size: 17px; border-radius: 12px; padding: 0.7rem 1.5rem; border: none;
        box-shadow: 0 6px 16px rgba(30, 58, 138, 0.3);
        transition: all 0.2s ease;
        font-family: 'B Nazanin', Tahoma, sans-serif !important;
    }}
    .stButton button:hover {{ transform: translateY(-2px); box-shadow: 0 8px 20px rgba(37, 99, 235, 0.4); }}
    </style>
""", unsafe_allow_html=True)

STATUS_COLORS = {
    'ثبت شده در سیستم': {'site_bg': '#fee2e2', 'site_fg': '#dc2626', 'excel_hex': 'FFCCCC', 'progress': 10},
    'برنامه‌ریزی‌شده': {'site_bg': '#fef3c7', 'site_fg': '#d97706', 'excel_hex': 'FFE599', 'progress': 25},
    'در حال تولید در سالن': {'site_bg': '#dbeafe', 'site_fg': '#2563eb', 'excel_hex': 'CCE5FF', 'progress': 50},
    'کسری تولید (نیازمند تکمیل)': {'site_bg': '#ffedd5', 'site_fg': '#c2410c', 'excel_hex': 'FCE4D6', 'progress': 60},
    'تکمیل تولید': {'site_bg': '#e0e7ff', 'site_fg': '#4338ca', 'excel_hex': 'C6C6FF', 'progress': 75},
    'کنترل کیفیت و بسته‌بندی': {'site_bg': '#fef9c3', 'site_fg': '#ca8a04', 'excel_hex': 'FFF3CD', 'progress': 85},
    'آماده تحویل در انبار': {'site_bg': '#dcfce7', 'site_fg': '#16a34a', 'excel_hex': 'D4EDDA', 'progress': 100},
    'ارسال شده به مقصد مشتری': {'site_bg': '#f3e8ff', 'site_fg': '#7e22ce', 'excel_hex': 'E2D9F3', 'progress': 100},
    'مرجوع شده': {'site_bg': '#ffe4e6', 'site_fg': '#9f1239', 'excel_hex': 'F8D7DA', 'progress': 0}
}

def log_activity(action_type, details):
    try:
        timestamp = jdatetime.datetime.now().strftime("%Y/%m/%d - %H:%M:%S")
        new_log = pd.DataFrame([{'تاریخ و ساعت': timestamp, 'نوع عملیات': action_type, 'شرح جزئیات': details}])
        if os.path.exists(AUDIT_LOG_FILE):
            df_log = pd.read_excel(AUDIT_LOG_FILE)
            df_log = pd.concat([new_log, df_log], ignore_index=True)
        else:
            df_log = new_log
        df_log.to_excel(AUDIT_LOG_FILE, index=False)
    except Exception:
        pass

def log_warehouse_txn(txn_type, code_val, prod_name, rolls_val, details_str):
    try:
        timestamp = jdatetime.datetime.now().strftime("%Y/%m/%d - %H:%M")
        new_w_log = pd.DataFrame([{
            'تاریخ و ساعت': timestamp, 'نوع تراکنش': txn_type, 'کد کالا': code_val, 'محصول': prod_name, 
            'تعداد طاقه': rolls_val, 'جزئیات': details_str
        }])
        if os.path.exists(WAREHOUSE_LOG_FILE):
            df_w = pd.read_excel(WAREHOUSE_LOG_FILE)
            df_w = pd.concat([new_w_log, df_w], ignore_index=True)
        else:
            df_w = new_w_log
        df_w.to_excel(WAREHOUSE_LOG_FILE, index=False)
    except Exception:
        pass

def update_excel_colors(file_path):
    try:
        if not os.path.exists(file_path): return
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        status_col_idx = None
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=col).value == 'وضعیت':
                status_col_idx = col
                break
        if status_col_idx:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=status_col_idx)
                val = cell.value
                if val in STATUS_COLORS:
                    hex_c = STATUS_COLORS[val]['excel_hex']
                    cell.fill = PatternFill(start_color=hex_c, end_color=hex_c, fill_type="solid")
                cell.alignment = Alignment(horizontal='center', vertical='center')
        wb.save(file_path)
    except Exception:
        pass

def get_users_db():
    if os.path.exists(USERS_DB_FILE):
        try:
            df_u = pd.read_excel(USERS_DB_FILE)
            if 'permissions' not in df_u.columns:
                df_u['permissions'] = "ثبت سفارش,کارتابل مشتریان,مدیریت سفارشات,برنامه‌ریزی تولید,تولید و کارگاه,انباردار,کنترل کیفیت,گزارشات کلیدی,مدیریت کاربران"
            return df_u
        except Exception:
            pass
    return pd.DataFrame([
        {"username": "admin", "password": "123", "role": "مدیرعامل", "name": "مدیرعامل محترم", "permissions": "ثبت سفارش,کارتابل مشتریان,مدیریت سفارشات,برنامه‌ریزی تولید,تولید و کارگاه,انباردار,کنترل کیفیت,گزارشات کلیدی,مدیریت کاربران"},
        {"username": "sales", "password": "123", "role": "مدیر فروش", "name": "آقای اسماعیلی (مسئول فروش)", "permissions": "ثبت سفارش,کارتابل مشتریان,مدیریت سفارشات"}
    ])

def get_products_master():
    if os.path.exists(PRODUCTS_MASTER_FILE):
        try:
            df_p = pd.read_excel(PRODUCTS_MASTER_FILE)
            if 'کد کالا' in df_p.columns and 'محصول / مدل فویل' in df_p.columns:
                df_p['کد کالا'] = df_p['کد کالا'].astype(str).str.strip()
                df_p['محصول / مدل فویل'] = df_p['محصول / مدل فویل'].astype(str).str.strip()
                return df_p
        except Exception:
            pass
    return pd.DataFrame({
        'کد کالا': ['2601/876', '1201/11', '1201/99'],
        'محصول / مدل فویل': ['فندقي قرمز مات', 'سفيد مات', 'مشکي مات']
    })

class PDF(FPDF):
    def header(self):
        self.set_fill_color(30, 58, 138)
        self.rect(0, 0, 210, 25, style='F')
        self.set_text_color(255, 255, 255)
        self.set_font('helvetica', 'B', 11)
        self.cell(0, 10, 'DIBAPOSH ESPADANA - MANAGEMENT REPORT', align='C', ln=True)
        self.ln(15)

    def footer(self):
        self.set_y(-15)
        self.set_font('helvetica', 'I', 8)
        self.set_text_color(128, 128, 128)
        self.cell(0, 10, f'Page {self.page_no()}', align='C')

def generate_customer_pdf(cust_name, orders_df):
    pdf = PDF(orientation='P', unit='mm', format='A4')
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    pdf.set_text_color(30, 58, 138)
    pdf.set_font('helvetica', 'B', 14)
    safe_cust_name = ''.join(c for c in str(cust_name) if ord(c) < 128)
    if not safe_cust_name: safe_cust_name = "Customer"
    
    pdf.cell(0, 8, f'Customer: {safe_cust_name}', ln=True, align='L')
    pdf.set_font('helvetica', '', 9)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 5, f'Report Date: {jdatetime.datetime.now().strftime("%Y/%m/%d - %H:%M")}', ln=True, align='L')
    pdf.ln(4)
    
    pdf.set_fill_color(37, 99, 235)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('helvetica', 'B', 8)
    
    headers = ['Order ID', 'Product Code', 'Product Model', 'Rolls', 'Status']
    widths = [20, 25, 65, 18, 42]
    
    for i, h in enumerate(headers):
        pdf.cell(widths[i], 7, h, 1, 0, 'C', fill=True)
    pdf.ln()
    
    pdf.set_font('helvetica', '', 8)
    pdf.set_text_color(30, 41, 59)
    
    for index, row in orders_df.iterrows():
        oid = ''.join(c for c in str(row['شماره سفارش']) if ord(c) < 128)
        if not oid: oid = "N/A"
        
        pcode = ''.join(c for c in str(row['کد کالا']) if ord(c) < 128)
        prod = ''.join(c for c in str(row['محصول / مدل فویل']) if ord(c) < 128)[:25]
        rolls = str(row['تعداد طاقه'])
        status = ''.join(c for c in str(row['وضعیت']) if ord(c) < 128)
        
        pdf.cell(widths[0], 6, oid, 1, 0, 'C')
        pdf.cell(widths[1], 6, pcode, 1, 0, 'C')
        pdf.cell(widths[2], 6, prod, 1, 0, 'L')
        pdf.cell(widths[3], 6, rolls, 1, 0, 'C')
        pdf.cell(widths[4], 6, status, 1, 0, 'C')
        pdf.ln()
        
    pdf.ln(4)
    total_rolls = orders_df['تعداد طاقه'].sum()
    pdf.cell(0, 7, f'Total Rolls: {total_rolls}', ln=True, align='R')

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
    pdf.output(temp_file.name)
    return temp_file.name

if 'authenticated' not in st.session_state:
    st.session_state.authenticated = False
if 'user_role' not in st.session_state:
    st.session_state.user_role = None
if 'username' not in st.session_state:
    st.session_state.username = None
if 'user_permissions' not in st.session_state:
    st.session_state.user_permissions = []

USERS_DF = get_users_db()

col_switch1, col_switch2 = st.columns([10, 2])
with col_switch2:
    mode_toggle = st.toggle("🌙 حالت تاریک", value=st.session_state.dark_mode)
    if mode_toggle != st.session_state.dark_mode:
        st.session_state.dark_mode = mode_toggle
        st.rerun()

if not st.session_state.authenticated:
    st.markdown("""
        <div class='landing-header'>
            <h1 style='font-weight: 900; font-size: 34px; margin-bottom: 8px;'>🧵 شرکت تولیدی و صنعتی دیبا پوش اسپادانا</h1>
            <p style='font-size: 18px; font-weight: 700;'>سامانه جامع فرماندهی صنعتی، برنامه‌ریزی تولید و انبارداری هوشمند (مبنا: طاقه ۵۰۰ متری)</p>
        </div>
    """, unsafe_allow_html=True)
    
    col_intro, col_form = st.columns([1.3, 1], gap="large")
    
    with col_intro:
        st.markdown("""
            <div class='company-intro-card'>
                <div style='font-size: 26px; font-weight: 900; margin-bottom: 15px; color: #1e3a8a;'>🚀 پورتال یکپارچه مدیریت صنعتی</div>
                <div style='font-size: 16px; font-weight: 700; line-height: 1.8; margin-bottom: 25px;'>
                    به سامانه تخصصی مدیریت تولید و انبارداری دیباپوش اسپادانا خوش آمدید. این پورتال جهت نظارت دقیق بر خطوط هات‌استمپ، ردیابی سفارشات، مدیریت انبار مستقل و کنترل کیفی پیشرفته طراحی شده است.
                </div>
                <div class='portal-badge'>📊 مدیریت هوشمند سفارشات و تخصیص انبار / تولید</div>
                <div class='portal-badge'>🏭 فرماندهی برنامه‌ریزی خط (MPS) بر اساس تفکیک کد و مدل کالا</div>
                <div class='portal-badge'>🛡️ ارزیابی تخصصی کیفیت (QC) و پایش لحظه‌ای پیشرفت</div>
            </div>
        """, unsafe_allow_html=True)
        
    with col_form:
        st.markdown("""
            <div class='login-form-card'>
                <div style='font-size: 22px; font-weight: 900; margin-bottom: 8px; color: #1e3a8a;'>ورود به حساب کاربری</div>
                <div style='font-size: 15px; font-weight: 700; margin-bottom: 25px;'>لطفاً جهت دسترسی به بخش سازمانی خود وارد شوید</div>
        """, unsafe_allow_html=True)
        
        with st.form("login_form"):
            u_input = st.text_input("نام کاربری (Username)")
            p_input = st.text_input("رمز عبور (Password)", type="password")
            login_btn = st.form_submit_button("ورود به سامانه 🚀", use_container_width=True)
            if login_btn:
                user_match = USERS_DF[USERS_DF['username'].astype(str).str.strip() == str(u_input).strip()]
                if not user_match.empty and str(user_match.iloc[0]['password']) == str(p_input):
                    st.session_state.authenticated = True
                    st.session_state.user_role = str(user_match.iloc[0]['role'])
                    st.session_state.username = str(user_match.iloc[0]['name'])
                    perms_val = str(user_match.iloc[0]['permissions'])
                    st.session_state.user_permissions = [p.strip() for p in perms_val.split(',') if p.strip()]
                    log_activity("ورود به سیستم", f"کاربر [{st.session_state.username}] وارد شد.")
                    st.success(f"خوش آمدید، {st.session_state.username}!")
                    st.rerun()
                else:
                    st.error("نام کاربری یا رمز عبور اشتباه است.")
        st.markdown("</div>", unsafe_allow_html=True)

else:
    st.markdown(f"""
        <div class='hero-container' style='padding: 25px; margin-bottom: 20px;'>
            <h2 style='color: white; margin: 0; font-weight: 900;'>شرکت تولیدی و صنعتی دیبا پوش اسپادانا</h2>
            <p style='color: #e2e8f0; margin: 5px 0 0 0; font-weight: 700;'>کاربر گرامی: <b>{st.session_state.username}</b> | سمت: <b>{st.session_state.user_role}</b> (مبنای طاقه ۵۰۰ متری)</p>
        </div>
    """, unsafe_allow_html=True)
    
    if st.button("🚪 خروج از حساب کاربری"):
        log_activity("خروج", f"کاربر [{st.session_state.username}] خارج شد.")
        st.session_state.authenticated = False
        st.session_state.user_role = None
        st.session_state.username = None
        st.session_state.user_permissions = []
        st.rerun()

    perms = st.session_state.user_permissions

    tab_names = []
    tab_mapping = {}

    if "ثبت سفارش" in perms:
        tab_names.append("🛒 ثبت سفارش")
        tab_mapping["🛒 ثبت سفارش"] = "tab1"
    if "کارتابل مشتریان" in perms:
        tab_names.append("👤 کارتابل مشتریان")
        tab_mapping["👤 کارتابل مشتریان"] = "tab2"
    if "مدیریت سفارشات" in perms:
        tab_names.append("📊 مدیریت سفارشات")
        tab_mapping["📊 مدیریت سفارشات"] = "tab3"
    if "برنامه‌ریزی تولید" in perms:
        tab_names.append("🧠 برنامه‌ریزی تولید (MPS)")
        tab_mapping["🧠 برنامه‌ریزی تولید (MPS)"] = "tab4"
    if "تولید و کارگاه" in perms:
        tab_names.append("🏭 تولید و کارگاه")
        tab_mapping["🏭 تولید و کارگاه"] = "tab5"
    if "انباردار" in perms:
        tab_names.append("📦 انباردار و واحد برش")
        tab_mapping["📦 انباردار و واحد برش"] = "tab6"
    if "کنترل کیفیت" in perms:
        tab_names.append("🔍 کنترل کیفیت (QC)")
        tab_mapping["🔍 کنترل کیفیت (QC)"] = "tab7"
    if "گزارشات کلیدی" in perms:
        tab_names.append("📈 گزارشات کلیدی")
        tab_mapping["📈 گزارشات کلیدی"] = "tab8"
    if "مدیریت کاربران" in perms:
        tab_names.append("👥 مدیریت کاربران")
        tab_mapping["👥 مدیریت کاربران"] = "tab9"

    if not tab_names:
        st.error("🚨 حساب کاربری شما هیچ دسترسی فعالی ندارد. لطفاً با مدیر سیستم تماس بگیرید.")
        st.stop()

    created_tabs = st.tabs(tab_names)
    tabs_dict = {tab_names[i]: created_tabs[i] for i in range(len(tab_names))}

    # ۱. ثبت سفارش
    if "🛒 ثبت سفارش" in tabs_dict:
        with tabs_dict["🛒 ثبت سفارش"]:
            st.subheader("🛒 ثبت سفارشات هوشمند (تفکیک دقیق کد کالا و نام محصول، کنترل انبار و تخصیص هوشمند موجودی قبلی)")
            
            with st.expander("➕ تعریف و ذخیره محصول / کد جدید در دیتابیس"):
                with st.form("quick_add_product_form"):
                    col_qp1, col_qp2 = st.columns(2)
                    with col_qp1:
                        new_c_code = st.text_input("کد کالا جدید (مثلاً 2601/876)", value="")
                    with col_qp2:
                        new_c_name = st.text_input("نام و مدل فویل جدید", value="")
                    
                    if st.form_submit_button("💾 ذخیره محصول جدید در دیتابیس"):
                        if new_c_code and new_c_name:
                            df_pm = get_products_master()
                            new_p_row = pd.DataFrame([{'کد کالا': str(new_c_code).strip(), 'محصول / مدل فویل': str(new_c_name).strip()}])
                            df_pm = pd.concat([df_pm, new_p_row], ignore_index=True)
                            df_pm.to_excel(PRODUCTS_MASTER_FILE, index=False)

                            if os.path.exists(INVENTORY_FILE):
                                df_im = pd.read_excel(INVENTORY_FILE)
                            else:
                                df_im = pd.DataFrame(columns=['کد کالا', 'محصول / مدل فویل', 'موجودی انبار (طاقه)', 'موجودی انبار (متر)', 'رول مادر قواره‌ای (متر)', 'سایز برش عرض (cm)', 'طاقه رزرو انبار', 'تاریخ انبار خوردن', 'مکان قفسه'])
                            
                            new_inv_row = pd.DataFrame([{
                                'کد کالا': str(new_c_code).strip(), 'محصول / مدل فویل': str(new_c_name).strip(),
                                'موجودی انبار (طاقه)': 0, 'موجودی انبار (متر)': 0.0, 'رول مادر قواره‌ای (متر)': 0.0,
                                'سایز برش عرض (cm)': 16.0, 'طاقه رزرو انبار': 0,
                                'تاریخ انبار خوردن': jdatetime.datetime.now().strftime("%Y/%m/%d"), 'مکان قفسه': 'Rack-1-Shelf1'
                            }])
                            df_im = pd.concat([df_im, new_inv_row], ignore_index=True)
                            df_im.to_excel(INVENTORY_FILE, index=False)

                            log_activity("تعریف محصول جدید", f"کد: {new_c_code} - نام: {new_c_name}")
                            st.success(f"محصول جدید [{new_c_name}] با کد [{new_c_code}] با موفقیت در دیتابیس ذخیره شد! 🎉")
                            st.rerun()
                        else:
                            st.error("لطفاً هم کد کالا و هم نام محصول را وارد کنید.")

            df_master = get_products_master()
            df_master['نمایش_کامل'] = df_master['محصول / مدل فویل'].astype(str) + " - " + df_master['کد کالا'].astype(str)
            product_options = df_master['نمایش_کامل'].tolist()

            df_inv_live = pd.read_excel(INVENTORY_FILE) if os.path.exists(INVENTORY_FILE) else pd.DataFrame()
            if not df_inv_live.empty and 'طاقه رزرو انبار' not in df_inv_live.columns:
                df_inv_live['طاقه رزرو انبار'] = 0
            
            col_s1, col_s2, col_s3 = st.columns(3)
            with col_s1:
                oid = st.text_input("شماره فاکتور / سفارش (مثلاً 1405-101)", key="cart_oid")
            with col_s2:
                cname = st.text_input("نام مشتری / شرکت (مثلاً آقای ایکس)", key="cart_cname")
            with col_s3:
                manual_order_date = st.text_input("تاریخ ثبت (YYYY/MM/DD)", value=jdatetime.datetime.now().strftime("%Y/%m/%d"))
                
            st.markdown("---")
            if 'cart_items' not in st.session_state: st.session_state.cart_items = []

            selected_display = st.selectbox("انتخاب محصول و کد کالا", product_options, key="cart_prod")
            
            selected_row = df_master[df_master['نمایش_کامل'] == selected_display].iloc[0]
            sel_code = str(selected_row['کد کالا']).strip()
            sel_name = str(selected_row['محصول / مدل فویل']).strip()

            current_inv_rolls = 0
            if not df_inv_live.empty:
                match_inv = df_inv_live[(df_inv_live['کد کالا'].astype(str).str.strip() == sel_code) | (df_inv_live['محصول / مدل فویل'].astype(str).str.strip() == sel_name)]
                if not match_inv.empty:
                    current_inv_rolls = int(match_inv.iloc[0]['موجودی انبار (طاقه)'])

            owner_info_str = "آزاد (بدون صاحب)"
            ready_in_warehouse = pd.DataFrame()
            if os.path.exists(ORDERS_FILE):
                df_ord_check_owner = pd.read_excel(ORDERS_FILE)
                ready_in_warehouse = df_ord_check_owner[(df_ord_check_owner['کد کالا'].astype(str).str.strip() == sel_code) & (df_ord_check_owner['وضعیت'] == 'آماده تحویل در انبار')]
                if not ready_in_warehouse.empty:
                    owners_list = []
                    for idx, r_item in ready_in_warehouse.iterrows():
                        owners_list.append(f"مشتری: {r_item['نام مشتری']} (فاکتور {r_item['شماره سفارش']} - {r_item['تعداد طاقه']} طاقه)")
                    owner_info_str = " | ".join(owners_list)

            if current_inv_rolls > 0:
                st.info(f"💡 **موجودی فعلی انبار برای این کالا:** {current_inv_rolls} طاقه — 🏷️ **مالکیت/تخصیص فعلی در انبار:** {owner_info_str}")
                steal_from_owner = st.checkbox("⚡ اختصاص موجودی انبارِ متعلق به سفارش قبلی به این سفارش فوری (سفارش قبلی دچار کسری شده و به صف برنامه‌ریزی تولید بازمی‌گردد)")
            else:
                st.warning(f"⚠️ **اخطار انبار:** موجودی این کالا در انبار **صفر** است! تمام درخواست‌های این کالا باید از طریق تولید تأمین شود.")

            rolls_in = st.number_input("تعداد طاقه کل درخواستی مشتری", min_value=1, value=3, step=1, key="cart_rolls_in")
            item_priority = st.selectbox("🎯 اولویت این قلم کالا", ["عادی", "اورژانسی ⚡", "نیازمند تسویه 💳"], key="cart_item_priority")
            
            st.markdown("##### ⚙️ نحوه تأمین و تخصیص:")
            col_t1, col_t2 = st.columns(2)
            with col_t1:
                max_allowed_wh = rolls_in if ('steal_from_owner' in locals() and steal_from_owner) else min(rolls_in, current_inv_rolls)
                from_warehouse_rolls = st.number_input("تعداد طاقه ارسالی از انبار", min_value=0, max_value=int(rolls_in), value=int(max_allowed_wh), step=1)
            with col_t2:
                from_production_rolls = rolls_in - from_warehouse_rolls
                st.metric("تعداد طاقه نیازمند تولید", f"{from_production_rolls} طاقه")

            warehouse_error = False
            if from_warehouse_rolls > current_inv_rolls and not ('steal_from_owner' in locals() and steal_from_owner):
                st.error(f"🚨 **خطای موجودی انبار:** تعداد درخواستی از انبار ({from_warehouse_rolls} طاقه) بیشتر از موجودی واقعی انبار ({current_inv_rolls} طاقه) است!")
                warehouse_error = True

            if st.button("➕ افزودن این قلم به سبد خرید فاکتور"):
                if warehouse_error:
                    st.error("❌ به دلیل کسری موجودی انبار، امکان افزودن این قلم وجود ندارد. لطفاً مقدار سهم انبار را کاهش دهید.")
                elif sel_name:
                    if from_warehouse_rolls == rolls_in:
                        fulfillment_mode = '🟢 تأمین کامل از انبار'
                    elif from_warehouse_rolls == 0:
                        fulfillment_mode = '🟡 ساخت کامل از تولید'
                    else:
                        fulfillment_mode = '🔵 تأمین ترکیبی (انبار + تولید)'

                    if 'steal_from_owner' in locals() and steal_from_owner and not ready_in_warehouse.empty:
                        target_victim = ready_in_warehouse.iloc[0]
                        v_idx = df_ord_check_owner[df_ord_check_owner['شناسه سطر'].astype(str) == str(target_victim['شناسه سطر'])].index[0]
                        
                        new_vic_rolls = int(target_victim['تعداد طاقه']) - from_warehouse_rolls
                        if new_vic_rolls > 0:
                            df_ord_check_owner.loc[v_idx, 'تعداد طاقه'] = new_vic_rolls
                            df_ord_check_owner.loc[v_idx, 'متراژ کل (متر)'] = float(new_vic_rolls * METERS_PER_ROLL)
                            df_ord_check_owner.loc[v_idx, 'وضعیت'] = 'ثبت شده در سیستم'
                            df_ord_check_owner.loc[v_idx, 'کسری برش (طاقه)'] = from_warehouse_rolls
                        else:
                            df_ord_check_owner.loc[v_idx, 'وضعیت'] = 'مرجوع شده'
                            df_ord_check_owner.loc[v_idx, 'علت مرجوعی'] = 'تخصیص موجودی انبار به سفارش فوری جدید'
                        
                        df_ord_check_owner.to_excel(ORDERS_FILE, index=False)
                        update_excel_colors(ORDERS_FILE)
                        log_activity("جابجایی موجودی انبار", f"موجودی از سفارش {target_victim['شماره سفارش']} به سفارش فوری {oid} تخصیص یافت.")

                    st.session_state.cart_items.append({
                        'code': sel_code, 'product': sel_name, 'rolls': rolls_in, 'meters_total': float(rolls_in * METERS_PER_ROLL),
                        'stock_rolls': from_warehouse_rolls, 'prod_rolls': from_production_rolls, 'fulfillment': fulfillment_mode,
                        'priority': item_priority
                    })
                    st.success(f"[{sel_name} - کد: {sel_code}] با اولویت [{item_priority}] به تعداد {rolls_in} طاقه به سبد افزوده شد.")

            if st.session_state.cart_items:
                st.markdown("#### 🛒 اقلام موجود در سبد خرید جاری:")
                st.dataframe(pd.DataFrame(st.session_state.cart_items), use_container_width=True, hide_index=True)

                if st.button("🚀 ثبت نهایی کل سبد خرید در سیستم"):
                    if oid and cname:
                        df = pd.read_excel(ORDERS_FILE) if os.path.exists(ORDERS_FILE) else pd.DataFrame()
                        current_time_suffix = jdatetime.datetime.now().strftime(" - %H:%M")
                        final_date_str = manual_order_date.strip() + current_time_suffix
                        
                        new_rows = []
                        for item in st.session_state.cart_items:
                            row_id = f"{oid}-{item['code']}"
                            new_rows.append({
                                'وضعیت': 'ثبت شده در سیستم', 'شناسه سطر': row_id, 'شماره سفارش': str(oid),
                                'نام مشتری': str(cname), 'کد کالا': item['code'], 'محصول / مدل فویل': item['product'], 
                                'تعداد طاقه': item['rolls'], 'متراژ کل (متر)': item['meters_total'], 
                                'سهم انبار (طاقه)': item['stock_rolls'], 'سهم تولید (طاقه)': item['prod_rolls'], 'نحوه تامین و ارسال': item['fulfillment'],
                                'اولویت': item['priority'], 'ثبت‌کننده': f"{st.session_state.username} ({st.session_state.user_role})",
                                'تاریخ ثبت': final_date_str, 'متراژ رول مادر تولیدی': 0.0, 
                                'تاریخ تولید': '-', 'تاریخ تکمیل تولید': '-', 'تاریخ ارسال': '-', 'تاریخ مرجوعی': '-', 'علت مرجوعی': '-', 'کسری برش (طاقه)': 0
                            })
                        
                        df_final = pd.concat([df, pd.DataFrame(new_rows)], ignore_index=True)
                        df_final.to_excel(ORDERS_FILE, index=False)
                        update_excel_colors(ORDERS_FILE)
                        log_activity("ثبت سفارش", f"فاکتور {oid} برای مشتری {cname} ثبت شد.")
                        st.session_state.cart_items = []
                        st.success("سفارش با موفقیت ثبت نهایی شد! 🎉")
                        st.rerun()
                    else:
                        st.error("لطفاً شماره فاکتور و نام مشتری را وارد کنید.")

    # ۲. کارتابل مشتریان
    if "👤 کارتابل مشتریان" in tabs_dict:
        with tabs_dict["👤 کارتابل مشتریان"]:
            st.subheader("👤 کارتابل جامع مشتریان (رهگیری لحظه‌ای پیشرفت، تولید، برش و تحویل‌های جزئی)")
            
            if os.path.exists(ORDERS_FILE):
                df_cust = pd.read_excel(ORDERS_FILE)
                if not df_cust.empty:
                    for cust in df_cust['نام مشتری'].astype(str).unique():
                        cust_orders = df_cust[df_cust['نام مشتری'].astype(str) == cust]
                        tot_cust_rolls = cust_orders['تعداد طاقه'].sum()
                        
                        progress_vals = []
                        for st_val in cust_orders['وضعیت']:
                            progress_vals.append(STATUS_COLORS.get(st_val, {}).get('progress', 10))
                        avg_progress = int(sum(progress_vals) / len(progress_vals)) if progress_vals else 0

                        with st.expander(f"📁 مشتری: {cust} | مجموع کل سفارشات: {tot_cust_rolls} طاقه | 🟢 درصد آماده بودن: {avg_progress}%"):
                            col_p1, col_p2 = st.columns([3, 1])
                            with col_p1:
                                st.progress(avg_progress / 100.0, text=f"وضعیت آمادگی سفارشات مشتری: {avg_progress}%")
                            with col_p2:
                                st.markdown(f"<h3 style='color: #2563eb; text-align: center; margin: 0; font-weight: 900;'>{avg_progress}% 📊</h3>", unsafe_allow_html=True)

                            pdf_path = generate_customer_pdf(cust, cust_orders)
                            with open(pdf_path, "rb") as pdf_file:
                                st.download_button("📥 دانلود PDF گزارش مشتری", data=pdf_file.read(), file_name=f"Report_{cust}.pdf", mime="application/pdf", key=f"pdf_{cust}")
                            st.dataframe(cust_orders, use_container_width=True, hide_index=True)

    # ۳. مدیریت سفارشات
    if "📊 مدیریت سفارشات" in tabs_dict:
        with tabs_dict["📊 مدیریت سفارشات"]:
            st.subheader("📊 مدیریت سفارشات (ثبت و اصلاح قطعی تاریخ‌ها، وضعیت و مدیریت حرفه‌ای)")
            
            if os.path.exists(ORDERS_FILE):
                df_o = pd.read_excel(ORDERS_FILE)
                if not df_o.empty:
                    st.markdown("##### 📌 آمار و تعداد سفارشات به تفکیک وضعیت:")
                    status_counts = df_o['وضعیت'].value_counts()
                    
                    cols_cnt = st.columns(len(STATUS_COLORS))
                    for idx, (st_name, st_info) in enumerate(STATUS_COLORS.items()):
                        cnt_val = status_counts.get(st_name, 0)
                        cols_cnt[idx].markdown(f"""
                            <div class="metric-card">
                                <div style="font-size: 14px; color: {subtext_color}; font-weight: 700;">{st_name}</div>
                                <div style="font-size: 22px; font-weight: 900; color: {st_info['site_fg']}; margin-top: 5px;">{cnt_val} قلم</div>
                            </div>
                        """, unsafe_allow_html=True)
                    
                    st.markdown("---")

            col_m1, col_m2, col_m3, col_m4 = st.columns([1.5, 1.5, 1.5, 1.5])
            with col_m1:
                filter_status_choice = st.selectbox("🎯 فیلتر وضعیت", ["همه وضعیت‌ها"] + list(STATUS_COLORS.keys()))
            with col_m2:
                sort_by_choice = st.selectbox("🔀 سورت جدول بر اساس", ["نام مشتری (الف تا ی)", "تاریخ ثبت (جدیدترین)", "تاریخ ثبت (قدیمی‌ترین)", "تعداد طاقه (بیشترین)"])
            with col_m3:
                date_filter_mode = st.selectbox("📅 فیلتر بازه تاریخ", ["همه زمان‌ها", "ماه شمسی خاص", "روز خاص"])
            with col_m4:
                search_cust_query = st.text_input("🔍 جستجوی سریع مشتری یا فاکتور", value="")

            if date_filter_mode == "ماه شمسی خاص":
                target_m_val = st.text_input("ماه (مثلاً 1405/06)", value=jdatetime.datetime.now().strftime("%Y/%m"))
            elif date_filter_mode == "روز خاص":
                target_d_val = st.text_input("روز (مثلاً 1405/06/01)", value=jdatetime.datetime.now().strftime("%Y/%m/%d"))

            if os.path.exists(ORDERS_FILE):
                df_o = pd.read_excel(ORDERS_FILE)
                if not df_o.empty:
                    if filter_status_choice != "همه وضعیت‌ها":
                        df_o = df_o[df_o['وضعیت'] == filter_status_choice]

                    if date_filter_mode == "ماه شمسی خاص":
                        df_o = df_o[df_o['تاریخ ثبت'].astype(str).str.contains(target_m_val)]
                    elif date_filter_mode == "روز خاص":
                        df_o = df_o[df_o['تاریخ ثبت'].astype(str).str.contains(target_d_val)]

                    if search_cust_query:
                        mask = df_o['نام مشتری'].astype(str).str.contains(search_cust_query, case=False, na=False) | \
                               df_o['شماره سفارش'].astype(str).str.contains(search_cust_query, case=False, na=False)
                        df_o = df_o[mask]

                    if sort_by_choice == "نام مشتری (الف تا ی)":
                        df_o = df_o.sort_values(by=['نام مشتری', 'تاریخ ثبت'], ascending=[True, False]).reset_index(drop=True)
                    elif sort_by_choice == "تاریخ ثبت (جدیدترین)":
                        df_o = df_o.sort_values(by='تاریخ ثبت', ascending=False).reset_index(drop=True)
                    elif sort_by_choice == "تاریخ ثبت (قدیمی‌ترین)":
                        df_o = df_o.sort_values(by='تاریخ ثبت', ascending=True).reset_index(drop=True)
                    elif sort_by_choice == "تعداد طاقه (بیشترین)":
                        df_o = df_o.sort_values(by='تعداد طاقه', ascending=False).reset_index(drop=True)

                    def color_site_status(val):
                        if val in STATUS_COLORS:
                            c = STATUS_COLORS[val]
                            return f"background-color: {c['site_bg']}; color: {c['site_fg']}; font-weight: bold; border-radius: 4px;"
                        return ''

                    st.dataframe(df_o.style.map(color_site_status, subset=['وضعیت']), use_container_width=True, hide_index=True)
                    
                    st.markdown("---")
                    st.markdown("### ⚙️ پنل حرفه‌ای و جمع‌وجور مدیریت (ویرایش، اصلاح تاریخ‌ها و حذف)")
                    
                    row_ids = df_o['شناسه سطر'].astype(str).tolist() if 'شناسه سطر' in df_o.columns else df_o['شماره سفارش'].astype(str).tolist()
                    if row_ids:
                        selected_row_id = st.selectbox("انتخاب شناسه سطر مورد نظر جهت مدیریت", row_ids)
                        target_row = df_o[df_o['شناسه سطر'].astype(str) == selected_row_id]
                        if not target_row.empty:
                            r_idx = target_row.index[0]
                            
                            with st.form(f"form_manage_{selected_row_id}"):
                                st.markdown(f"🛠️ **ابزار مدیریت و اصلاح سفارش: [{selected_row_id}]**")
                                col_act1, col_act2 = st.columns(2)
                                with col_act1:
                                    edit_status = st.selectbox("تغییر وضعیت", list(STATUS_COLORS.keys()), index=list(STATUS_COLORS.keys()).index(df_o.loc[r_idx, 'وضعیت']) if df_o.loc[r_idx, 'وضعیت'] in STATUS_COLORS else 0)
                                    edit_rolls = st.number_input("تعداد طاقه", value=int(df_o.loc[r_idx, 'تعداد طاقه']), step=1)
                                    edit_date_reg = st.text_input("تاریخ ثبت", value=str(df_o.loc[r_idx, 'تاریخ ثبت']))
                                    edit_date_prod = st.text_input("تاریخ شروع تولید", value=str(df_o.loc[r_idx, 'تاریخ تولید']))
                                with col_act2:
                                    edit_date_comp = st.text_input("تاریخ تکمیل تولید", value=str(df_o.loc[r_idx, 'تاریخ تکمیل تولید']))
                                    edit_date_ship = st.text_input("تاریخ ارسال", value=str(df_o.loc[r_idx, 'تاریخ ارسال']))
                                    edit_date_ret = st.text_input("تاریخ مرجوعی", value=str(df_o.loc[r_idx, 'تاریخ مرجوعی']))
                                    edit_ret_reason = st.text_input("علت مرجوعی", value=str(df_o.loc[r_idx, 'علت مرجوعی']))

                                col_btn1, col_btn2 = st.columns([2, 1])
                                with col_btn1:
                                    save_changes = st.form_submit_button("💾 اعمال و ذخیره تغییرات و تاریخ‌ها")
                                with col_btn2:
                                    delete_order_btn = st.form_submit_button("🗑️ حذف سفارش")

                                if save_changes:
                                    current_today = jdatetime.datetime.now().strftime("%Y/%m/%d")
                                    df_o.loc[r_idx, 'وضعیت'] = edit_status
                                    df_o.loc[r_idx, 'تعداد طاقه'] = int(edit_rolls)
                                    df_o.loc[r_idx, 'متراژ کل (متر)'] = int(edit_rolls * METERS_PER_ROLL)
                                    df_o.loc[r_idx, 'تاریخ ثبت'] = edit_date_reg
                                    df_o.loc[r_idx, 'تاریخ تکمیل تولید'] = edit_date_comp
                                    df_o.loc[r_idx, 'علت مرجوعی'] = edit_ret_reason if edit_ret_reason else '-'
                                    
                                    if edit_status == 'در حال تولید در سالن' and (str(edit_date_prod) == '-' or not str(edit_date_prod)):
                                        df_o.loc[r_idx, 'تاریخ تولید'] = current_today
                                    else:
                                        df_o.loc[r_idx, 'تاریخ تولید'] = edit_date_prod

                                    if edit_status == 'ارسال شده به مقصد مشتری' and (str(edit_date_ship) == '-' or not str(edit_date_ship)):
                                        df_o.loc[r_idx, 'تاریخ ارسال'] = current_today
                                    else:
                                        df_o.loc[r_idx, 'تاریخ ارسال'] = edit_date_ship

                                    if edit_status == 'مرجوع شده' and (str(edit_date_ret) == '-' or not str(edit_date_ret)):
                                        df_o.loc[r_idx, 'تاریخ مرجوعی'] = current_today
                                    else:
                                        df_o.loc[r_idx, 'تاریخ مرجوعی'] = edit_date_ret

                                    df_o.to_excel(ORDERS_FILE, index=False)
                                    update_excel_colors(ORDERS_FILE)
                                    log_activity("مدیریت سفارش", f"سفارش {selected_row_id} به‌روزرسانی شد.")
                                    st.success("تغییرات و تاریخ‌ها با موفقیت ذخیره شد! ✅")
                                    st.rerun()

                                if delete_order_btn:
                                    df_o = df_o[df_o['شناسه سطر'].astype(str) != selected_row_id]
                                    df_o.to_excel(ORDERS_FILE, index=False)
                                    update_excel_colors(ORDERS_FILE)
                                    log_activity("حذف سفارش", f"سفارش {selected_row_id} حذف شد.")
                                    st.warning("سفارش حذف شد! 🗑️")
                                    st.rerun()
                else:
                    st.info("هیچ سفارشی در سیستم ثبت نشده است.")

    # ۴. برنامه‌ریزی تولید (MPS)
    if "🧠 برنامه‌ریزی تولید (MPS)" in tabs_dict:
        with tabs_dict["🧠 برنامه‌ریزی تولید (MPS)"]:
            st.subheader("🧠 برنامه‌ریزی تولید (MPS) - فیلتر وضعیت، جدول جامع و جداول مجزای هفتگی (شنبه تا پنج‌شنبه)")
            
            col_r1, col_r2 = st.columns([8, 2])
            with col_r2:
                if st.button("🔄 ریست کامل جداول این هفته"):
                    if os.path.exists(MPS_FILE):
                        pd.DataFrame(columns=[
                            'شناسه سطر', 'تاریخ برنامه‌ریزی تولید', 'روز هفته', 'شماره سفارش', 'نام مشتری', 
                            'کد کالا', 'محصول / مدل فویل', 'دستگاه تولید', 'طیف رنگی', 'تعداد طاقه تولید', 'متراژ تولید (متر)', 'اولویت', 'تاریخ ثبت سفارش', 'وضعیت اجرای خط', 'یادداشت برنامه‌ریز'
                        ]).to_excel(MPS_FILE, index=False)
                        st.success("جداول هفتگی از نو ریست شد! ✅")
                        st.rerun()

            df_ord_mps_raw = pd.read_excel(ORDERS_FILE) if os.path.exists(ORDERS_FILE) else pd.DataFrame()
            df_mps_existing = pd.read_excel(MPS_FILE) if os.path.exists(MPS_FILE) else pd.DataFrame()

            st.markdown("##### 📌 ۱. جدول سفارشات (با تفکیک دقیق ثبت شده، برنامه‌ریزی‌شده و کسری‌ها):")
            if not df_ord_mps_raw.empty:
                col_f1, col_f2, col_f3, col_f4 = st.columns(4)
                with col_f1:
                    status_filter_view = st.selectbox("📂 فیلتر وضعیت جدول", ["همه موارد", "سفارشات جدید (ثبت‌شده)", "برنامه‌ریزی‌شده‌ها", "در حال تولید در سالن", "کسری تولید (نیازمند تکمیل)", "تکمیل تولیدی‌ها"])
                with col_f2:
                    filter_reg_date = st.selectbox("📅 تاریخ ثبت", ["همه تاریخ‌ها"] + df_ord_mps_raw['تاریخ ثبت'].astype(str).unique().tolist())
                with col_f3:
                    filter_urgency = st.selectbox("⚡ فوریت", ["همه موارد", "عادی", "اورژانسی ⚡", "نیازمند تسویه 💳"])
                with col_f4:
                    filter_need = st.selectbox("⚙️ نوع تامین", ["همه موارد", "🟡 ساخت کامل از تولید", "🔵 تأمین ترکیبی"])

                df_filtered_table = df_ord_mps_raw.copy()
                
                if status_filter_view == "سفارشات جدید (ثبت‌شده)":
                    df_filtered_table = df_filtered_table[df_filtered_table['وضعیت'] == 'ثبت شده در سیستم']
                elif status_filter_view == "برنامه‌ریزی‌شده‌ها":
                    df_filtered_table = df_filtered_table[df_filtered_table['وضعیت'] == 'برنامه‌ریزی‌شده']
                elif status_filter_view == "در حال تولید در سالن":
                    df_filtered_table = df_filtered_table[df_filtered_table['وضعیت'] == 'در حال تولید در سالن']
                elif status_filter_view == "کسری تولید (نیازمند تکمیل)":
                    df_filtered_table = df_filtered_table[df_filtered_table['وضعیت'] == 'کسری تولید (نیازمند تکمیل)']
                elif status_filter_view == "تکمیل تولیدی‌ها":
                    df_filtered_table = df_filtered_table[df_filtered_table['وضعیت'].isin(['تکمیل تولید', 'کنترل کیفیت و بسته‌بندی', 'آماده تحویل در انبار'])]

                if filter_reg_date != "همه تاریخ‌ها":
                    df_filtered_table = df_filtered_table[df_filtered_table['تاریخ ثبت'].astype(str) == filter_reg_date]
                if filter_urgency != "همه موارد":
                    df_filtered_table = df_filtered_table[df_filtered_table['اولویت'].astype(str) == filter_urgency]
                if filter_need != "همه موارد":
                    df_filtered_table = df_filtered_table[df_filtered_table['نحوه تامین و ارسال'].astype(str).str.contains(filter_need[:10])]

                st.dataframe(df_filtered_table[['شماره سفارش', 'تاریخ ثبت', 'نام مشتری', 'کد کالا', 'محصول / مدل فویل', 'تعداد طاقه', 'اولویت', 'نحوه تامین و ارسال', 'وضعیت']], use_container_width=True, hide_index=True)
                
                st.markdown("---")
                st.markdown("##### 🎯 تخصیص سفارش به روزهای هفته (برنامه‌ریزی تولید و کسری‌ها):")
                
                week_days_full = ["شنبه", "یکشنبه", "دوشنبه", "سه‌شنبه", "چهارشنبه", "پنج‌شنبه"]
                
                with st.form("schedule_to_weekly_days_form"):
                    row_ids = df_filtered_table[df_filtered_table['وضعیت'].isin(['ثبت شده در سیستم', 'کسری تولید (نیازمند تکمیل)'])].sort_values(by='اولویت', ascending=False)['شناسه سطر'].astype(str).tolist() if not df_filtered_table.empty else []
                    if row_ids:
                        sel_order_id = st.selectbox("انتخاب سفارش جهت زمان‌بندی (جدید یا کسری)", row_ids)
                        target_ord_info = df_filtered_table[df_filtered_table['شناسه سطر'].astype(str) == sel_order_id].iloc[0]
                        
                        col_w1, col_w2, col_w3 = st.columns(3)
                        with col_w1:
                            chosen_weekday = st.selectbox("انتخاب روز هفته", week_days_full)
                        with col_w2:
                            today_j = jdatetime.datetime.now()
                            base_saturday = today_j - jdatetime.timedelta(days=today_j.weekday() if today_j.weekday() != 6 else 6)
                            calculated_date = (base_saturday + jdatetime.timedelta(days=week_days_full.index(chosen_weekday))).strftime("%Y/%m/%d")
                            chosen_date_w = st.text_input("تاریخ تولید (هماهنگ با روز هفته)", value=calculated_date)
                        with col_w3:
                            chosen_mach_w = st.selectbox("انتخاب دستگاه", [
                                "دستگاه چاپ هات‌استمپ 1 (Rotogravure)", 
                                "دستگاه چاپ هات‌استمپ 2 (UV Flatbed)"
                            ])
                        
                        note_w = st.text_input("یادداشت برنامه‌ریز", value=f"برنامه‌ریزی شد برای روز {chosen_weekday}.")
                        
                        if st.form_submit_button("🚀 ثبت و انتقال برنامه به جدول روزانه"):
                            new_w_row = pd.DataFrame([{
                                'شناسه سطر': target_ord_info['شناسه سطر'], 'تاریخ برنامه‌ریزی تولید': chosen_date_w, 
                                'روز هفته': chosen_weekday, 'شماره سفارش': target_ord_info['شماره سفارش'], 
                                'تاریخ ثبت سفارش': target_ord_info['تاریخ ثبت'], 'نام مشتری': target_ord_info['نام مشتری'], 
                                'کد کالا': target_ord_info['کد کالا'], 'محصول / مدل فویل': target_ord_info['محصول / مدل فویل'], 'دستگاه تولید': chosen_mach_w, 
                                'طیف رنگی': f"{chosen_weekday} - طیف استاندارد", 'تعداد طاقه تولید': target_ord_info['سهم تولید (طاقه)'], 
                                'متراژ تولید (متر)': float(target_ord_info['سهم تولید (طاقه)'] * METERS_PER_ROLL),
                                'اولویت': target_ord_info['اولویت'], 'وضعیت اجرای خط': 'برنامه‌ریزی شده', 'یادداشت برنامه‌ریز': note_w
                            }])
                            
                            if os.path.exists(MPS_FILE):
                                df_old_mps = pd.read_excel(MPS_FILE)
                                if 'روز هفته' not in df_old_mps.columns: df_old_mps['روز هفته'] = 'شنبه'
                                df_old_mps = df_old_mps[df_old_mps['شناسه سطر'].astype(str) != str(target_ord_info['شناسه سطر'])]
                                df_final_mps = pd.concat([df_old_mps, new_w_row], ignore_index=True)
                            else:
                                df_final_mps = new_w_row
                                
                            df_final_mps.to_excel(MPS_FILE, index=False)
                            
                            df_ord_mps_raw.loc[df_ord_mps_raw['شناسه سطر'].astype(str) == sel_order_id, 'وضعیت'] = 'برنامه‌ریزی‌شده'
                            df_ord_mps_raw.to_excel(ORDERS_FILE, index=False)
                            update_excel_colors(ORDERS_FILE)
                            
                            st.success(f"سفارش با موفقیت برنامه‌ریزی شد و به جدول روز **{chosen_weekday}** منتقل گردید! ✅")
                            st.rerun()
                    else:
                        st.info("هیچ سفارش جدید یا کسری‌داری برای برنامه‌ریزی در این فیلتر وجود ندارد.")
            else:
                st.info("هیچ سفارشی ثبت نشده است.")

            st.markdown("---")
            st.markdown("### 📅 ۶ جدول مجزا و پشت‌سرهم برای روزهای هفته (از شنبه تا پنج‌شنبه):")
            
            week_days_full = ["شنبه", "یکشنبه", "دوشنبه", "سه‌شنبه", "چهارشنبه", "پنج‌شنبه"]
            
            for day_name in week_days_full:
                st.markdown(f"#### 🗓️ جدول روز: **{day_name}**")
                if not df_mps_existing.empty and 'روز هفته' in df_mps_existing.columns:
                    df_day_sub = df_mps_existing[df_mps_existing['روز هفته'].astype(str) == day_name]
                    if not df_day_sub.empty:
                        st.dataframe(df_day_sub[['شماره سفارش', 'تاریخ برنامه‌ریزی تولید', 'نام مشتری', 'کد کالا', 'محصول / مدل فویل', 'تعداد طاقه تولید', 'دستگاه تولید', 'اولویت', 'یادداشت برنامه‌ریز']], use_container_width=True, hide_index=True)
                    else:
                        st.info(f"هیچ برنامه‌ای در جدول روز {day_name} نیست.")
                else:
                    st.info(f"هیچ برنامه‌ای ثبت نشده است.")
                st.markdown("---")

    # ۵. تولید و کارگاه
    if "🏭 تولید و کارگاه" in tabs_dict:
        with tabs_dict["🏭 تولید و کارگاه"]:
            st.subheader("🏭 تولید و کارگاه - کارتابل خلوت و منظم سرپرست خط (تکالیف امروز، معوقات و رول‌های فعال)")
            
            df_mps_ws_chk = pd.read_excel(MPS_FILE) if os.path.exists(MPS_FILE) else pd.DataFrame()
            today_date_str = jdatetime.datetime.now().strftime("%Y/%m/%d")

            if not df_mps_ws_chk.empty and 'تاریخ برنامه‌ریزی تولید' in df_mps_ws_chk.columns:
                df_approved_tasks = df_mps_ws_chk[df_mps_ws_chk['وضعیت اجرای خط'] == 'برنامه‌ریزی شده']
                df_overdue = df_approved_tasks[df_approved_tasks['تاریخ برنامه‌ریزی تولید'].astype(str) < today_date_str]
                df_today = df_approved_tasks[df_approved_tasks['تاریخ برنامه‌ریزی تولید'].astype(str) == today_date_str]

                if not df_overdue.empty:
                    st.markdown(f"""
                        <div class="notification-banner">
                            <span>🚨 اخطار معوقات کارگاه: تعداد <b>{len(df_overdue)} دستور کار از روزهای قبل</b> هنوز تکمیل نشده‌اند!</span>
                            <span style="background: rgba(255,255,255,0.2); padding: 4px 10px; border-radius: 8px; font-size: 14px;">اورژانسی</span>
                        </div>
                    """, unsafe_allow_html=True)

            wk_tab_today, wk_tab_overdue, wk_tab_running, wk_tab_done = st.tabs([
                "📅 ۱. تکالیف تولید امروز",
                "⏰ ۲. معوقات روزهای قبل (اورژانسی)",
                "⚡ ۳. ثبت رول مادر و اعلام کسری متریال",
                "✅ ۴. تکمیل‌شده‌ها (ارسال‌شده به انبار)"
            ])

            with wk_tab_today:
                st.markdown(f"##### 📅 دستورات کاری ابلاغ‌شده برای تولید امروز ({today_date_str}):")
                if not df_mps_ws_chk.empty and 'df_today' in locals() and not df_today.empty:
                    st.dataframe(df_today[['شماره سفارش', 'روز هفته', 'تاریخ برنامه‌ریزی تولید', 'نام مشتری', 'کد کالا', 'محصول / مدل فویل', 'تعداد طاقه تولید', 'دستگاه تولید', 'اولویت']], use_container_width=True, hide_index=True)
                else:
                    st.info("هیچ دستور کاری برای امروز ابلاغ نشده است.")

            with wk_tab_overdue:
                st.markdown("##### ⏰ دستورات کاری معوقه از روزهای قبل:")
                if not df_mps_ws_chk.empty and 'df_overdue' in locals() and not df_overdue.empty:
                    st.dataframe(df_overdue[['شماره سفارش', 'روز هفته', 'تاریخ برنامه‌ریزی تولید', 'نام مشتری', 'کد کالا', 'محصول / مدل فویل', 'تعداد طاقه تولید', 'دستگاه تولید', 'اولویت']], use_container_width=True, hide_index=True)
                else:
                    st.success("هیچ معوقه‌ای وجود ندارد؛ کارگاه به‌روز است! 🎉")

            with wk_tab_running:
                st.markdown("##### ⚡ ثبت متراژ واقعی رول مادر و بررسی دقیق کسری / مازاد تولید:")
                df_ws_all = pd.read_excel(MPS_FILE) if os.path.exists(MPS_FILE) else pd.DataFrame()
                
                ws_m1, ws_m2 = st.tabs(["🖨️ ثبت رول مادر در دستگاه 1", "⚡ ثبت رول مادر در دستگاه 2"])
                
                with ws_m1:
                    df_ws1 = df_ws_all[df_ws_all['دستگاه تولید'].astype(str).str.contains('Rotogravure') & (df_ws_all['وضعیت اجرای خط'].isin(['برنامه‌ریزی شده', 'در حال تولید']))] if not df_ws_all.empty else pd.DataFrame()
                    if not df_ws1.empty:
                        with st.form("ws1_form"):
                            s1_ids = df_ws1['شناسه سطر'].astype(str).tolist()
                            sel_s1 = st.selectbox("انتخاب دستور کار دستگاه 1", s1_ids, key="s1_k")
                            idx1 = df_ws_all[df_ws_all['شناسه سطر'].astype(str) == sel_s1].index[0]
                            
                            s1_status = st.selectbox("وضعیت اجرای خط", ['در حال تولید', 'تکمیل تولید'], index=0, key="st1_k")
                            mother_roll_meters = st.number_input("متراژ واقعی رول مادر تولید شده در شیفت (متر) *اجباری*", min_value=0.0, value=0.0, step=10.0, key="raw_m1")
                            
                            needed_r = int(df_ws_all.loc[idx1, 'تعداد طاقه تولید'])
                            needed_meters = needed_r * METERS_PER_ROLL

                            if mother_roll_meters > 0:
                                produced_rolls_equiv = int(mother_roll_meters // METERS_PER_ROLL)
                                deficit_rolls = max(0, needed_r - produced_rolls_equiv)
                                deficit_meters = max(0.0, needed_meters - mother_roll_meters)
                                surplus_meters = max(0.0, mother_roll_meters - needed_meters)

                                if mother_roll_meters < needed_meters:
                                    st.error(f"⚠️ **تولید ناقص (کسری متریال):** متراژ تولیدی ({mother_roll_meters} متر) کمتر از مقدار سفارش ({needed_meters} متر) است! مقدار **{deficit_meters} متر ({deficit_rolls} طاقه) کسری** دارد و این سفارش **تکمیل نشده** است.")
                                elif mother_roll_meters > needed_meters:
                                    st.success(f"🟢 **مازاد تولید:** متراژ تولیدی بیشتر از سفارش است و مقدار **{surplus_meters} متر مازاد** ثبت شد.")
                                else:
                                    st.success(f"✅ این رول مادر پوشش کامل و دقیق سفارش را می‌دهد.")

                            s1_date = st.text_input("تاریخ تولید (YYYY/MM/DD)", value=jdatetime.datetime.now().strftime("%Y/%m/%d"), key="d1_k")
                            s1_note = st.text_input("گزارش شیفت / توضیحات", value=str(df_ws_all.loc[idx1, 'یادداشت برنامه‌ریز']), key="n1_k")
                            
                            if st.form_submit_button("🚀 ثبت گزارش و ارسال به انبار"):
                                if mother_roll_meters <= 0:
                                    st.error("❌ پر کردن فیلد «متراژ واقعی رول مادر تولید شده» اجباری است و نمی‌تواند صفر باشد!")
                                else:
                                    df_ws_all.loc[idx1, 'یادداشت برنامه‌ریز'] = s1_note
                                    df_ws_all.to_excel(MPS_FILE, index=False)
                                    
                                    if os.path.exists(ORDERS_FILE):
                                        df_o_up = pd.read_excel(ORDERS_FILE)
                                        if sel_s1 in df_o_up['شناسه سطر'].astype(str).values:
                                            o_idx = df_o_up[df_o_up['شناسه سطر'].astype(str) == sel_s1].index[0]
                                            
                                            df_o_up.loc[o_idx, 'متراژ رول مادر تولیدی'] = float(mother_roll_meters)
                                            df_o_up.loc[o_idx, 'تاریخ تولید'] = s1_date
                                            df_o_up.loc[o_idx, 'کسری برش (طاقه)'] = deficit_rolls

                                            if mother_roll_meters < needed_meters:
                                                df_o_up.loc[o_idx, 'وضعیت'] = 'کسری تولید (نیازمند تکمیل)'
                                                df_o_up.loc[o_idx, 'تاریخ تکمیل تولید'] = '-'
                                            else:
                                                df_o_up.loc[o_idx, 'وضعیت'] = 'تکمیل تولید'
                                                df_o_up.loc[o_idx, 'تاریخ تکمیل تولید'] = s1_date

                                            df_o_up.to_excel(ORDERS_FILE, index=False)
                                            update_excel_colors(ORDERS_FILE)

                                    log_activity("کارگاه دستگاه 1", f"رول مادر با متراژ {mother_roll_meters} متر ثبت شد.")
                                    st.success("گزارش شیفت ثبت شد! ✅")
                                    st.rerun()
                    else:
                        st.info("هیچ دستور کاری برای دستگاه ۱ نیست.")

                with ws_m2:
                    df_ws2 = df_ws_all[df_ws_all['دستگاه تولید'].astype(str).str.contains('Flatbed') & (df_ws_all['وضعیت اجرای خط'].isin(['برنامه‌ریزی شده', 'در حال تولید']))] if not df_ws_all.empty else pd.DataFrame()
                    if not df_ws2.empty:
                        with st.form("ws2_form"):
                            s2_ids = df_ws2['شناسه سطر'].astype(str).tolist()
                            sel_s2 = st.selectbox("انتخاب دستور کار دستگاه 2", s2_ids, key="s2_k")
                            idx2 = df_ws_all[df_ws_all['شناسه سطر'].astype(str) == sel_s2].index[0]
                            
                            s2_status = st.selectbox("وضعیت اجرای خط", ['در حال تولید', 'تکمیل تولید'], index=0, key="st2_k")
                            mother_roll_meters2 = st.number_input("متراژ واقعی رول مادر تولید شده در شیفت (متر) *اجباری*", min_value=0.0, value=0.0, step=10.0, key="raw_m2")
                            
                            needed_r2 = int(df_ws_all.loc[idx2, 'تعداد طاقه تولید'])
                            needed_meters2 = needed_r2 * METERS_PER_ROLL

                            if mother_roll_meters2 > 0:
                                produced_rolls_equiv2 = int(mother_roll_meters2 // METERS_PER_ROLL)
                                deficit_rolls2 = max(0, needed_r2 - produced_rolls_equiv2)
                                deficit_meters2 = max(0.0, needed_meters2 - mother_roll_meters2)

                                if mother_roll_meters2 < needed_meters2:
                                    st.error(f"⚠️ **تولید ناقص (کسری متریال):** متراژ تولیدی ({mother_roll_meters2} متر) کمتر از سفارش است! مقدار **{deficit_meters2} متر ({deficit_rolls2} طاقه) کسری** دارد و تکمیل نشده است.")
                                else:
                                    st.success(f"✅ این رول مادر پوشش کامل سفارش را می‌دهد.")

                            s2_date = st.text_input("تاریخ تولید (YYYY/MM/DD)", value=jdatetime.datetime.now().strftime("%Y/%m/%d"), key="d2_k")
                            s2_note = st.text_input("گزارش شیفت / توضیحات", value=str(df_ws_all.loc[idx2, 'یادداشت برنامه‌ریز']), key="n2_k")
                            
                            if st.form_submit_button("🚀 ثبت گزارش و ارسال به انبار"):
                                if mother_roll_meters2 <= 0:
                                    st.error("❌ پر کردن فیلد «متراژ واقعی رول مادر تولید شده» اجباری است و نمی‌تواند صفر باشد!")
                                else:
                                    df_ws_all.loc[idx2, 'یادداشت برنامه‌ریز'] = s2_note
                                    df_ws_all.to_excel(MPS_FILE, index=False)
                                    
                                    if os.path.exists(ORDERS_FILE):
                                        df_o_up = pd.read_excel(ORDERS_FILE)
                                        if sel_s2 in df_o_up['شناسه سطر'].astype(str).values:
                                            o_idx2 = df_o_up[df_o_up['شناسه سطر'].astype(str) == sel_s2].index[0]
                                            
                                            df_o_up.loc[o_idx2, 'متراژ رول مادر تولیدی'] = float(mother_roll_meters2)
                                            df_o_up.loc[o_idx2, 'تاریخ تولید'] = s2_date
                                            df_o_up.loc[o_idx2, 'کسری برش (طاقه)'] = deficit_rolls2

                                            if mother_roll_meters2 < needed_meters2:
                                                df_o_up.loc[o_idx2, 'وضعیت'] = 'کسری تولید (نیازمند تکمیل)'
                                                df_o_up.loc[o_idx2, 'تاریخ تکمیل تولید'] = '-'
                                            else:
                                                df_o_up.loc[o_idx2, 'وضعیت'] = 'تکمیل تولید'
                                                df_o_up.loc[o_idx2, 'تاریخ تکمیل تولید'] = s2_date

                                            df_o_up.to_excel(ORDERS_FILE, index=False)
                                            update_excel_colors(ORDERS_FILE)

                                    log_activity("کارگاه دستگاه 2", f"رول مادر با متراژ {mother_roll_meters2} متر ثبت شد.")
                                    st.success("گزارش شیفت ثبت شد! ✅")
                                    st.rerun()
                    else:
                        st.info("هیچ دستور کاری برای دستگاه ۲ نیست.")

            with wk_tab_done:
                st.markdown("##### ✅ تاریخچه رول‌های مادری که در کارگاه کامل تولید شده‌اند:")
                if not df_ws_all.empty:
                    df_wk_done = df_ws_all[df_ws_all['وضعیت اجرای خط'] == 'تکمیل تولید']
                    if not df_wk_done.empty:
                        st.dataframe(df_wk_done[['شماره سفارش', 'نام مشتری', 'کد کالا', 'محصول / مدل فویل', 'تعداد طاقه تولید', 'متراژ تولید (متر)', 'دستگاه تولید']], use_container_width=True, hide_index=True)
                    else:
                        st.info("هیچ پارت تکمیل‌شده‌ای در این بخش وجود ندارد.")

    # ۶. انباردار و واحد برش
    if "📦 انباردار و واحد برش" in tabs_dict:
        with tabs_dict["📦 انباردار و واحد برش"]:
            st.subheader("📦 انباردار و واحد برش - تفکیک بخش ورود/خروج دستی و میز اسلیت پیشرفته")
            
            df_wh_notif = pd.read_excel(ORDERS_FILE) if os.path.exists(ORDERS_FILE) else pd.DataFrame()
            if not df_wh_notif.empty:
                pending_wh_count = len(df_wh_notif[df_wh_notif['وضعیت'].isin(['تکمیل تولید', 'کنترل کیفیت و بسته‌بندی'])])
                if pending_wh_count > 0:
                    st.markdown(f"""
                        <div class="warehouse-alert">
                            <span>📦 آلارم انبار: تعداد <b>{pending_wh_count} رول مادر کامل</b> از کارگاه به انبار رسیده و منتظر اسلیت و برش است!</span>
                            <span style="background: rgba(255,255,255,0.2); padding: 4px 10px; border-radius: 8px; font-size: 14px;">واحد اسلیت</span>
                        </div>
                    """, unsafe_allow_html=True)

            wh_tab_manual, wh_tab_cutting, wh_tab_stats, wh_tab_manager_log = st.tabs([
                "📥 ۱. ورود و خروج دستی انبار",
                "✂️ ۲. میز برش، اسلیت و تخصیص سهم مشتری",
                "📊 ۳. آمار لحظه‌ای برش و پیگیری سهم مشتری",
                "👑 ۴. پنل مدیریت و ویرایش اطلاعات ثبت‌شده (اصلاح خطا)"
            ])

            try:
                df_inv = pd.read_excel(INVENTORY_FILE) if os.path.exists(INVENTORY_FILE) else pd.DataFrame()
            except Exception:
                df_inv = pd.DataFrame()

            if df_inv.empty:
                df_inv = pd.DataFrame(columns=['کد کالا', 'محصول / مدل فویل', 'موجودی انبار (طاقه)', 'موجودی انبار (متر)', 'رول مادر قواره‌ای (متر)', 'سایز برش عرض (cm)', 'طاقه رزرو انبار', 'تاریخ انبار خوردن', 'مکان قفسه'])

            if 'رول مادر قواره‌ای (متر)' not in df_inv.columns: df_inv['رول مادر قواره‌ای (متر)'] = 0.0
            if 'طاقه رزرو انبار' not in df_inv.columns: df_inv['طاقه رزرو انبار'] = 0
            if 'سایز برش عرض (cm)' not in df_inv.columns: df_inv['سایز برش عرض (cm)'] = 16.0
            df_inv.to_excel(INVENTORY_FILE, index=False)

            if os.path.exists(ORDERS_FILE):
                try:
                    df_ord_inv = pd.read_excel(ORDERS_FILE)
                    active_orders = df_ord_inv[~df_ord_inv['وضعیت'].isin(['ارسال شده به مقصد مشتری', 'مرجوع شده'])]
                    if not active_orders.empty and 'سهم انبار (طاقه)' in active_orders.columns:
                        reserved_df = active_orders.groupby('محصول / مدل فویل')['سهم انبار (طاقه)'].sum().reset_index()
                        reserved_df.rename(columns={'سهم انبار (طاقه)': 'طاقه رزرو انبار'}, inplace=True)
                        df_inv = df_inv.merge(reserved_df, on='محصول / مدل فویل', how='left')
                        df_inv['طاقه رزرو انبار'] = df_inv['طاقه رزرو انبار'].fillna(0).astype(int)
                except Exception:
                    pass

            if 'طاقه رزرو انبار' not in df_inv.columns:
                df_inv['طاقه رزرو انبار'] = 0

            df_inv['کسری انبار (طاقه)'] = df_inv['طاقه رزرو انبار'] - df_inv['موجودی انبار (طاقه)']
            df_inv['کسری انبار (طاقه)'] = df_inv['کسری انبار (طاقه)'].apply(lambda x: max(0, x))

            with wh_tab_manual:
                st.markdown("##### 📥 ثبت ورود یا خروج دستی اقلام به انبار مستقل از خط تولید:")
                col_wh_op1, col_wh_op2 = st.columns(2)
                
                with col_wh_op1:
                    st.markdown("###### ➕ ورود موجودی طاقه به انبار")
                    with st.form("update_inv_form"):
                        prod_list_inv = df_inv['محصول / مدل فویل'].tolist() if not df_inv.empty else []
                        sel_p = st.selectbox("انتخاب محصول", prod_list_inv)
                        add_rolls = st.number_input("اضافه کردن تعداد طاقه", min_value=0, value=2, step=1)
                        cut_cm_val = st.number_input("سایز برش عرض (cm)", min_value=5.0, value=16.0, step=1.0)
                        inv_date_in = st.text_input("تاریخ انبار خوردن (YYYY/MM/DD)", value=jdatetime.datetime.now().strftime("%Y/%m/%d"))
                        new_bin = st.text_input("مکان قفسه", value="Rack-1-Shelf1")
                        if st.form_submit_button("ثبت ورود موجودی ➕"):
                            if not df_inv.empty and sel_p in df_inv['محصول / مدل فویل'].values:
                                r_idx = df_inv[df_inv['محصول / مدل فویل'] == sel_p].index[0]
                                c_code_val = str(df_inv.loc[r_idx, 'کد کالا'])
                                curr_rolls = int(df_inv.loc[r_idx, 'موجودی انبار (طاقه)'])
                                new_rolls_val = curr_rolls + add_rolls
                                df_inv.loc[r_idx, 'موجودی انبار (طاقه)'] = new_rolls_val
                                df_inv.loc[r_idx, 'موجودی انبار (متر)'] = float(new_rolls_val * METERS_PER_ROLL)
                                df_inv.loc[r_idx, 'سایز برش عرض (cm)'] = float(cut_cm_val)
                                df_inv.loc[r_idx, 'تاریخ انبار خوردن'] = inv_date_in
                                if new_bin: df_inv.loc[r_idx, 'مکان قفسه'] = new_bin
                                df_inv.to_excel(INVENTORY_FILE, index=False)
                                log_warehouse_txn("ورودی دستی انبار", c_code_val, sel_p, add_rolls, f"ورود دستی - عرض برش: {cut_cm_val}cm")
                                st.success("موجودی انبار با موفقیت به‌روز شد! ✅")
                                st.rerun()

                with col_wh_op2:
                    st.markdown("###### 🆕 تعریف کالای جدید در انبار")
                    with st.form("add_new_prod_form"):
                        new_code = st.text_input("کد کالا", value="999")
                        new_prod_name = st.text_input("نام و مدل فویل جدید")
                        init_rolls = st.number_input("موجودی اولیه (طاقه)", min_value=0, value=2, step=1)
                        cut_cm = st.number_input("سایز برش عرض (cm)", min_value=5.0, value=16.0, step=1.0)
                        inv_date_new = st.text_input("تاریخ انبار خوردن", value=jdatetime.datetime.now().strftime("%Y/%m/%d"))
                        new_shelf = st.text_input("مکان قفسه", value="Rack-5-Shelf1")
                        if st.form_submit_button("افزودن محصول جدید 🌟"):
                            if new_prod_name:
                                new_row = pd.DataFrame([{
                                    'کد کالا': str(new_code).strip(),
                                    'محصول / مدل فویل': str(new_prod_name).strip(),
                                    'موجودی انبار (طاقه)': init_rolls,
                                    'موجودی انبار (متر)': float(init_rolls * METERS_PER_ROLL),
                                    'رول مادر قواره‌ای (متر)': 0.0,
                                    'سایز برش عرض (cm)': cut_cm,
                                    'تاریخ انبار خوردن': inv_date_new,
                                    'مکان قفسه': new_shelf,
                                    'طاقه رزرو انبار': 0
                                }])
                                df_inv = pd.concat([df_inv, new_row], ignore_index=True)
                                df_inv.to_excel(INVENTORY_FILE, index=False)
                                log_warehouse_txn("تعریف کالا", str(new_code), str(new_prod_name), init_rolls, f"تعریف اولیه کالا")
                                st.success("محصول جدید در انبار ثبت شد! 🎉")
                                st.rerun()

            with wh_tab_cutting:
                st.markdown("##### ✂️ میز برش و اسلیت رول مادر (محاسبه دقیق کسری و مازاد طاقه‌ها):")
                if os.path.exists(ORDERS_FILE):
                    df_wh_auto = pd.read_excel(ORDERS_FILE)
                    pending_workshop = df_wh_auto[df_wh_auto['وضعیت'].isin(['تکمیل تولید', 'کنترل کیفیت و بسته‌بندی'])]
                    
                    if not pending_workshop.empty:
                        w_row_ids = pending_workshop['شناسه سطر'].astype(str).tolist()
                        sel_w_id = st.selectbox("انتخاب پارت تولیدی رول مادر جهت اسلیت در انبار", w_row_ids, key="wh_workshop_sel")
                        
                        w_idx = df_wh_auto[df_wh_auto['شناسه سطر'].astype(str) == sel_w_id].index[0]
                        w_cust = df_wh_auto.loc[w_idx, 'نام مشتری']
                        w_prod = df_wh_auto.loc[w_idx, 'محصول / مدل فویل']
                        w_needed_rolls = int(df_wh_auto.loc[w_idx, 'تعداد طاقه'])
                        w_mother_meters = float(df_wh_auto.loc[w_idx, 'متراژ رول مادر تولیدی']) if 'متراژ رول مادر تولیدی' in df_wh_auto.columns else 0.0

                        st.info(f"مشتری: **{w_cust}** | محصول: **{w_prod}** | طاقه درخواستی فاکتور: **{w_needed_rolls} طاقه** | رول مادر تولید شده: **{w_mother_meters} متر**")

                        with st.form("warehouse_cutting_form"):
                            col_c1, col_c2 = st.columns(2)
                            with col_c1:
                                assigned_rolls = st.number_input("تعداد طاقه ۵۰۰ متری تحویلی در این مرحله (سهم مشتری)", min_value=0, value=min(w_needed_rolls, int(w_mother_meters // METERS_PER_ROLL)), step=1)
                            with col_c2:
                                cut_width_cm = st.number_input("عرض برش نهایی طاقه‌ها (cm)", min_value=5.0, value=16.0, step=1.0)
                            
                            used_meters = assigned_rolls * METERS_PER_ROLL
                            remainder_meters = max(0.0, w_mother_meters - used_meters)
                            
                            progress_percent = int((assigned_rolls / w_needed_rolls) * 100) if w_needed_rolls > 0 else 0
                            st.progress(min(1.0, progress_percent / 100.0), text=f"درصد تکمیل و اسلیت این سفارش: {progress_percent}%")
                            
                            current_deficit = max(0, w_needed_rolls - assigned_rolls)
                            surplus_rolls = max(0, assigned_rolls - w_needed_rolls)
                            
                            if current_deficit > 0:
                                st.warning(f"⚠️ **کسری در واحد برش:** این فاکتور نیاز به {w_needed_rolls} طاقه داشت اما {assigned_rolls} طاقه برش خورد؛ تعداد **{current_deficit} طاقه کسری** ثبت شد.")
                            elif surplus_rolls > 0:
                                st.success(f"🟢 **مازاد در برش:** تعداد {surplus_rolls} طاقه مازاد بر سفارش به انبار اضافه شد.")

                            slit_action = st.selectbox("وضعیت نهایی واحد برش در این مرحله", [
                                "اسلیت کامل، تخصیص به مشتری و آماده تحویل 🟢", 
                                "اسلیت مرحله‌ای (نیمه‌کاره / ادامه در روزهای بعد) 🟡", 
                                "دارای نقص یا ضایعات در حین برش ❌"
                            ])
                            warehouse_fault_note = st.text_input("توضیحات انبار", value="بدون نقص.")
                            
                            submit_wh_cut = st.form_submit_button("🛡️ تایید نهایی اسلیت و اعمال به انبار و سهم مشتری")
                            if submit_wh_cut:
                                current_today = jdatetime.datetime.now().strftime("%Y/%m/%d")
                                
                                if "نقص" in slit_action:
                                    df_wh_auto.loc[w_idx, 'وضعیت'] = 'مرجوع شده'
                                    df_wh_auto.loc[w_idx, 'علت مرجوعی'] = f"نقص واحد برش: {warehouse_fault_note}"
                                    df_wh_auto.loc[w_idx, 'تاریخ مرجوعی'] = current_today
                                else:
                                    if assigned_rolls >= w_needed_rolls:
                                        df_wh_auto.loc[w_idx, 'وضعیت'] = 'آماده تحویل در انبار'
                                    else:
                                        df_wh_auto.loc[w_idx, 'وضعیت'] = 'تکمیل تولید'
                                
                                df_wh_auto.loc[w_idx, 'کسری برش (طاقه)'] = current_deficit
                                
                                if not df_inv.empty and w_prod in df_inv['محصول / مدل فویل'].values:
                                    inv_idx = df_inv[df_inv['محصول / مدل فویل'] == w_prod].index[0]
                                    c_code_str = str(df_inv.loc[inv_idx, 'کد کالا'])
                                    curr_r = int(df_inv.loc[inv_idx, 'موجودی انبار (طاقه)'])
                                    curr_q = float(df_inv.loc[inv_idx, 'رول مادر قواره‌ای (متر)']) if 'رول مادر قواره‌ای (متر)' in df_inv.columns else 0.0
                                    
                                    new_r = curr_r + min(assigned_rolls, w_needed_rolls)
                                    new_q = curr_q + remainder_meters + (surplus_rolls * METERS_PER_ROLL)
                                    
                                    df_inv.loc[inv_idx, 'موجودی انبار (طاقه)'] = new_r
                                    df_inv.loc[inv_idx, 'موجودی انبار (متر)'] = float(new_r * METERS_PER_ROLL)
                                    df_inv.loc[inv_idx, 'رول مادر قواره‌ای (متر)'] = float(new_q)
                                    df_inv.loc[inv_idx, 'سایز برش عرض (cm)'] = float(cut_width_cm)
                                    df_inv.to_excel(INVENTORY_FILE, index=False)
                                    log_warehouse_txn("اسلیت رول مادر", c_code_str, w_prod, assigned_rolls, f"مشتری: {w_cust} - سهم تحویلی: {assigned_rolls} طاقه")

                                df_wh_auto.to_excel(ORDERS_FILE, index=False)
                                update_excel_colors(ORDERS_FILE)
                                log_activity("واحد برش و انبار", f"پارت {sel_w_id} اسلیت شد. طاقه‌ها: {assigned_rolls}")
                                st.success(f"عملیات اسلیت با موفقیت ثبت شد! ✅")
                                st.rerun()
                    else:
                        st.info("هیچ رول مادری آماده‌ای در صف انتظار انبار برای اسلیت وجود ندارد.")

            with wh_tab_stats:
                st.markdown("##### 📊 آمار لحظه‌ای برش، مقدار سهم مشتری و وضعیت تکمیل:")
                if os.path.exists(ORDERS_FILE):
                    df_stat = pd.read_excel(ORDERS_FILE)
                    if not df_stat.empty:
                        st.dataframe(df_stat[['شماره سفارش', 'نام مشتری', 'کد کالا', 'محصول / مدل فویل', 'تعداد طاقه', 'وضعیت', 'متراژ رول مادر تولیدی', 'تاریخ تکمیل تولید', 'کسری برش (طاقه)']], use_container_width=True, hide_index=True)
                    else:
                        st.info("هیچ داده‌ای ثبت نشده است.")

            with wh_tab_manager_log:
                st.markdown("👑 **پنل مدیریت و ویرایش اطلاعات ثبت‌شده (اصلاح خطا):**")
                if os.path.exists(ORDERS_FILE):
                    df_all_edits = pd.read_excel(ORDERS_FILE)
                    if not df_all_edits.empty:
                        edit_row_id = st.selectbox("انتخاب شناسه سفارش جهت اصلاح یا ویرایش اطلاعات", df_all_edits['شناسه سطر'].astype(str).tolist(), key="edit_err_row")
                        row_to_edit_idx = df_all_edits[df_all_edits['شناسه سطر'].astype(str) == edit_row_id].index[0]
                        
                        with st.form("fix_operator_error_form"):
                            f_cust = st.text_input("نام مشتری", value=str(df_all_edits.loc[row_to_edit_idx, 'نام مشتری']))
                            f_rolls = st.number_input("تعداد طاقه", value=int(df_all_edits.loc[row_to_edit_idx, 'تعداد طاقه']), step=1)
                            f_status = st.selectbox("وضعیت سفارش", list(STATUS_COLORS.keys()), index=list(STATUS_COLORS.keys()).index(df_all_edits.loc[row_to_edit_idx, 'وضعیت']) if df_all_edits.loc[row_to_edit_idx, 'وضعیت'] in STATUS_COLORS else 0)
                            
                            if st.form_submit_button("💾 ذخیره و اصلاح اطلاعات"):
                                df_all_edits.loc[row_to_edit_idx, 'نام مشتری'] = f_cust
                                df_all_edits.loc[row_to_edit_idx, 'تعداد طاقه'] = int(f_rolls)
                                df_all_edits.loc[row_to_edit_idx, 'متراژ کل (متر)'] = int(f_rolls * METERS_PER_ROLL)
                                df_all_edits.loc[row_to_edit_idx, 'وضعیت'] = f_status
                                df_all_edits.to_excel(ORDERS_FILE, index=False)
                                update_excel_colors(ORDERS_FILE)
                                log_activity("اصلاح اطلاعات", f"سفارش {edit_row_id} توسط مدیریت ویرایش شد.")
                                st.success("اطلاعات با موفقیت اصلاح شد! ✅")
                                st.rerun()

            st.markdown("---")
            st.markdown("##### 📋 جدول موجودی فیزیکی و رول‌های قواره‌ای انبار:")
            if not df_inv.empty:
                st.dataframe(df_inv[['کد کالا', 'محصول / مدل فویل', 'موجودی انبار (طاقه)', 'رول مادر قواره‌ای (متر)', 'سایز برش عرض (cm)', 'مکان قفسه', 'طاقه رزرو انبار']], use_container_width=True, hide_index=True)

    # ۷. کنترل کیفیت (QC)
    if "🔍 کنترل کیفیت (QC)" in tabs_dict:
        with tabs_dict["🔍 کنترل کیفیت (QC)"]:
            st.subheader("🔍 بخش تخصصی کنترل کیفیت (QC) - تفکیک پارت‌ها به همراه تاریخ و مشتری و آرشیو خودکار")
            
            if os.path.exists(ORDERS_FILE):
                df_ord_qc = pd.read_excel(ORDERS_FILE)
                qc_done_df = pd.read_excel(QC_FILE) if os.path.exists(QC_FILE) else pd.DataFrame()
                already_qc_ids = qc_done_df['شناسه سطر'].astype(str).tolist() if not qc_done_df.empty and 'شناسه سطر' in qc_done_df.columns else []
                
                qc_candidates = df_ord_qc[(df_ord_qc['وضعیت'].isin(['تکمیل تولید', 'کنترل کیفیت و بسته‌بندی'])) & (~df_ord_qc['شناسه سطر'].astype(str).isin(already_qc_ids))]
                
                if not qc_candidates.empty:
                    st.markdown("##### 📌 انتخاب پارت تولیدی جهت بازرسی و ارزیابی کیفیت:")
                    
                    qc_candidates['نمایش_گزینه'] = qc_candidates['شماره سفارش'].astype(str) + " | " + qc_candidates['نام مشتری'].astype(str) + " | [کد: " + qc_candidates['کد کالا'].astype(str) + "] " + qc_candidates['محصول / مدل فویل'].astype(str) + " (تاریخ تولید: " + qc_candidates['تاریخ تولید'].astype(str) + ")"
                    
                    qc_options = qc_candidates['نمایش_گزینه'].tolist()
                    sel_qc_display = st.selectbox("انتخاب پارت تولیدی", qc_options)
                    
                    selected_qc_row = qc_candidates[qc_candidates['نمایش_گزینه'] == sel_qc_display].iloc[0]
                    sel_qc_id = str(selected_qc_row['شناسه سطر'])
                    q_cust = selected_qc_row['نام مشتری']
                    q_code = selected_qc_row['کد کالا']
                    q_prod = selected_qc_row['محصول / مدل فویل']
                    q_rolls = selected_qc_row['تعداد طاقه']
                    q_date_prod = selected_qc_row['تاریخ تولید']
                    
                    st.info(f"مشتری: **{q_cust}** | محصول: [کد: **{q_code}**] **{q_prod}** | حجم پارت: **{q_rolls} طاقه** | تاریخ تولید: **{q_date_prod}**")

                    with st.form("qc_inspection_form"):
                        st.markdown("##### 🧪 ارزیابی پارامترهای فنی فویل هات‌استمپ:")
                        
                        col_qc1, col_qc2 = st.columns(2)
                        with col_qc1:
                            glue_quality = st.selectbox("کیفیت چسب و میزان چسبندگی (Adhesion)", [
                                "عالی و استاندارد (بدون ریزش)", 
                                "متوسط نیازمند بررسی لایه زیرین", 
                                "ضعیف / عدم چسبندگی مناسب ❌"
                            ])
                            hotstamp_quality = st.selectbox("کیفیت هات‌استمپ شدن (فویل‌کوبی و وضوح)", [
                                "شفاف، دقیق و بدون نقص", 
                                "قابل قبول با خطای جزیی لبه‌ها", 
                                "دارای زدگی یا ریزش طرح ❌"
                            ])
                        with col_qc2:
                            rub_resistance = st.selectbox("مقاومت در برابر سایش و خراش (Rub Test)", [
                                "مقاوم و تثبیت‌شده", 
                                "متوسط", 
                                "خش‌پذیر و نامقاوم ❌"
                            ])
                            visual_appearance = st.selectbox("ارزیابی ظاهری (رنگ، براقیت و حباب)", [
                                "بدون حباب، خط و خش یا ایراد رنگی", 
                                "دارای نقص جزئی ظاهری", 
                                "مردود ظاهری / سایه‌دار ❌"
                            ])

                        qc_final_verdict = st.selectbox("نتیجه نهایی ارزیابی QC", [
                            "تایید شده جهت بسته‌بندی و انبار 🟢", 
                            "نیازمند اصلاح یا بازکاری در سالن 🟡", 
                            "مردود قطعی / مرجوعی 🔴"
                        ])
                        
                        qc_inspector_name = st.text_input("نام بازرس / مسئول کنترل کیفیت", value=st.session_state.username)
                        qc_date_in = st.text_input("تاریخ ارزیابی QC (YYYY/MM/DD)", value=jdatetime.datetime.now().strftime("%Y/%m/%d"))
                        qc_notes = st.text_input("توضیحات تکمیلی و جمع‌بندی بازرس", value="ارزیابی فنی پارت تولیدی با موفقیت انجام شد.")

                        submit_qc_btn = st.form_submit_button("🛡️ تایید و آرشیو نهایی در دیتابیس")
                        if submit_qc_btn:
                            if "تایید" in qc_final_verdict:
                                df_ord_qc.loc[df_ord_qc['شناسه سطر'].astype(str) == sel_qc_id, 'وضعیت'] = 'آماده تحویل در انبار'
                            else:
                                df_ord_qc.loc[df_ord_qc['شناسه سطر'].astype(str) == sel_qc_id, 'وضعیت'] = 'مرجوع شده'
                                df_ord_qc.loc[df_ord_qc['شناسه سطر'].astype(str) == sel_qc_id, 'علت مرجوعی'] = f"مردود QC: {qc_notes}"

                            df_ord_qc.to_excel(ORDERS_FILE, index=False)
                            update_excel_colors(ORDERS_FILE)

                            new_qc_record = pd.DataFrame([{
                                'شناسه سطر': sel_qc_id, 'شماره سفارش': selected_qc_row['شماره سفارش'],
                                'نام مشتری': q_cust, 'کد کالا': q_code, 'محصول / مدل فویل': q_prod, 'تعداد طاقه': q_rolls, 'تاریخ تولید': q_date_prod,
                                'کیفیت چسب و چسبندگی': glue_quality, 'کیفیت هات‌استمپ شدن': hotstamp_quality,
                                'مقاومت در برابر سایش': rub_resistance, 'ارزیابی ظاهری': visual_appearance,
                                'نتیجه ارزیابی QC': qc_final_verdict, 'تاریخ ارزیابی': qc_date_in,
                                'مسئول کنترل کیفی': qc_inspector_name, 'توضیحات بازرس': qc_notes
                            }])

                            if os.path.exists(QC_FILE):
                                df_qc_old = pd.read_excel(QC_FILE)
                                df_qc_final = pd.concat([df_qc_old, new_qc_record], ignore_index=True)
                            else:
                                df_qc_final = new_qc_record

                            df_qc_final.to_excel(QC_FILE, index=False)
                            log_activity("کنترل کیفیت QC", f"پارت {sel_qc_id} ({q_prod}) بازرسی و در دیتابیس آرشیو شد.")
                            st.success("کنترل کیفیت تایید شد و این محصول از صفحه خارج و در دیتابیس آرشیو گردید! ✅")
                            st.rerun()
                else:
                    st.info("هیچ پارت تولیدی جدیدی نیازمند کنترل کیفیت در این بخش نیست (همه موارد بازرسی و آرشیو شده‌اند).")

                st.markdown("---")
                st.markdown("##### 📋 آرشیو کامل دیتابیس کنترل کیفیت (QC):")
                if os.path.exists(QC_FILE):
                    st.dataframe(pd.read_excel(QC_FILE), use_container_width=True, hide_index=True)

    # ۸. گزارشات کلیدی
    if "📈 گزارشات کلیدی" in tabs_dict:
        with tabs_dict["📈 گزارشات کلیدی"]:
            st.subheader("📈 داشبورد جامع فرماندهی مدیریتی و گزارشات تحلیلی (کل، ماه، روز)")
            
            if os.path.exists(ORDERS_FILE):
                df_kpi = pd.read_excel(ORDERS_FILE)
                
                col_k1, col_k2, col_k3, col_k4 = st.columns(4)
                
                with col_k1:
                    st.markdown(f"""
                        <div class="metric-card">
                            <div style="font-size: 16px; color: {subtext_color}; font-weight: 900; margin-bottom: 8px;">کل سفارشات ثبتی</div>
                            <div style="font-size: 30px; font-weight: 900; color: #2563eb;">{len(df_kpi)}</div>
                        </div>
                    """, unsafe_allow_html=True)
                    
                with col_k2:
                    tot_rolls_val = f"{df_kpi['تعداد طاقه'].sum():,} طاقه" if not df_kpi.empty else "0 طاقه"
                    st.markdown(f"""
                        <div class="metric-card">
                            <div style="font-size: 16px; color: {subtext_color}; font-weight: 900; margin-bottom: 8px;">مجموع کل طاقه‌ها</div>
                            <div style="font-size: 30px; font-weight: 900; color: #16a34a;">{tot_rolls_val}</div>
                        </div>
                    """, unsafe_allow_html=True)
                    
                with col_k3:
                    tot_meters_val = f"{df_kpi['متراژ کل (متر)'].sum():,.0f} متر" if not df_kpi.empty else "0 متر"
                    st.markdown(f"""
                        <div class="metric-card">
                            <div style="font-size: 16px; color: {subtext_color}; font-weight: 900; margin-bottom: 8px;">مجموع متراژ فویل</div>
                            <div style="font-size: 30px; font-weight: 900; color: #d97706;">{tot_meters_val}</div>
                        </div>
                    """, unsafe_allow_html=True)
                    
                with col_k4:
                    active_count = len(df_kpi[~df_kpi['وضعیت'].isin(['ارسال شده به مقصد مشتری', 'مرجوع شده'])]) if not df_kpi.empty else 0
                    st.markdown(f"""
                        <div class="metric-card">
                            <div style="font-size: 16px; color: {subtext_color}; font-weight: 900; margin-bottom: 8px;">فاکتورهای فعال جریان</div>
                            <div style="font-size: 30px; font-weight: 900; color: #7e22ce;">{active_count}</div>
                        </div>
                    """, unsafe_allow_html=True)
                
                st.markdown("---")
                st.markdown("#### 🏆 تحلیل هوشمند پرفروش‌ترین رنگ‌ها / مدل‌ها و مشتریان برتر")
                
                rep_time_mode = st.radio("انتخاب بازه زمانی تحلیل:", ["کل تاریخچه", "ماه شمسی خاص", "روز خاص"], horizontal=True)
                
                df_filtered_rep = df_kpi.copy()
                if rep_time_mode == "ماه شمسی خاص":
                    r_month_in = st.text_input("ورود ماه (مثلاً 1405/06)", value=jdatetime.datetime.now().strftime("%Y/%m"))
                    df_filtered_rep = df_filtered_rep[df_filtered_rep['تاریخ ثبت'].astype(str).str.contains(r_month_in)]
                elif rep_time_mode == "روز خاص":
                    r_day_in = st.text_input("ورود روز (مثلاً 1405/06/01)", value=jdatetime.datetime.now().strftime("%Y/%m/%d"))
                    df_filtered_rep = df_filtered_rep[df_filtered_rep['تاریخ ثبت'].astype(str).str.contains(r_day_in)]

                if not df_filtered_rep.empty:
                    col_rep1, col_rep2 = st.columns(2)
                    color_palette = ['#1e3a8a', '#2563eb', '#3b82f6', '#d97706', '#16a34a', '#7e22ce', '#dc2626', '#0891b2']
                    
                    with col_rep1:
                        st.markdown("##### 🎨 سهم مدل‌ها و رنگ‌های پرفروش فویل:")
                        top_products = df_filtered_rep.groupby('محصول / مدل فویل')['تعداد طاقه'].sum().reset_index()
                        fig_prod = px.pie(top_products, names='محصول / مدل فویل', values='تعداد طاقه', hole=0.4, color_discrete_sequence=color_palette)
                        fig_prod.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color=text_color, family='B Nazanin, Tahoma'))
                        st.plotly_chart(fig_prod, use_container_width=True, config={'displayModeBar': False, 'staticPlot': True})
                        
                    with col_rep2:
                        st.markdown("##### 👑 سهم برترین مشتریان از کل فروش:")
                        top_customers = df_filtered_rep.groupby('نام مشتری')['تعداد طاقه'].sum().reset_index()
                        fig_cust = px.pie(top_customers, names='نام مشتری', values='تعداد طاقه', hole=0.4, color_discrete_sequence=color_palette)
                        fig_cust.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color=text_color, family='B Nazanin, Tahoma'))
                        st.plotly_chart(fig_cust, use_container_width=True, config={'displayModeBar': False, 'staticPlot': True})
                else:
                    st.info("هیچ داده‌ای در این بازه زمانی برای تحلیل وجود ندارد.")

                st.markdown("---")
                col_ch1, col_ch2 = st.columns(2)
                with col_ch1:
                    st.markdown("##### 📊 نمودار دایره‌ای وضعیت سفارشات:")
                    if not df_kpi.empty:
                        status_dist = df_kpi['وضعیت'].value_counts().reset_index()
                        status_dist.columns = ['وضعیت', 'تعداد']
                        fig_status = px.pie(status_dist, names='وضعیت', values='تعداد', hole=0.4, color_discrete_sequence=color_palette)
                        fig_status.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color=text_color, family='B Nazanin, Tahoma'))
                        st.plotly_chart(fig_status, use_container_width=True, config={'displayModeBar': False, 'staticPlot': True})
                
                with col_ch2:
                    st.markdown("##### ⚖️ سهم روش تأمین (انبار در برابر تولید):")
                    if not df_kpi.empty and 'سهم انبار (طاقه)' in df_kpi.columns:
                        sum_wh = df_kpi['سهم انبار (طاقه)'].sum()
                        sum_pr = df_kpi['سهم تولید (طاقه)'].sum()
                        df_fulfillment_pie = pd.DataFrame({'منبع تأمین': ['تأمین از انبار', 'تأمین از تولید'], 'تعداد طاقه': [sum_wh, sum_pr]})
                        fig_full = px.pie(df_fulfillment_pie, names='منبع تأمین', values='تعداد طاقه', hole=0.4, color_discrete_sequence=['#16a34a', '#2563eb'])
                        fig_full.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color=text_color, family='B Nazanin, Tahoma'))
                        st.plotly_chart(fig_full, use_container_width=True, config={'displayModeBar': False, 'staticPlot': True})

                st.markdown("---")
                st.markdown("#### 📦 وضعیت موجودی کل انبار و رزروها:")
                if os.path.exists(INVENTORY_FILE):
                    df_inv_rep = pd.read_excel(INVENTORY_FILE)
                    st.dataframe(df_inv_rep, use_container_width=True, hide_index=True)

                st.markdown("---")
                st.markdown("#### 🕵️‍♂️ تاریخچه کامل لاگ، رویدادها و بررسی دسترسی‌ها:")
                if os.path.exists(AUDIT_LOG_FILE):
                    st.dataframe(pd.read_excel(AUDIT_LOG_FILE), use_container_width=True, hide_index=True)

    # ۹. مدیریت کاربران
    if "👥 مدیریت کاربران" in tabs_dict:
        with tabs_dict["👥 مدیریت کاربران"]:
            st.subheader("👥 پنل مدیریت کاربران و دسترسی‌های ترکیبی چندگانه")
            
            st.markdown("##### ➕ تعریف یا ویرایش دسترسی‌های کاربر:")
            with st.form("manage_users_form"):
                u_username = st.text_input("نام کاربری انگلیسی (مثلاً esmaeili)")
                u_password = st.text_input("رمز عبور", value="123")
                u_name = st.text_input("نام و نام خانوادگی / سمت (مثلاً آقای اسماعیلی)")
                
                st.markdown("##### 🔑 انتخاب بخش‌های مجاز برای این کاربر (دسترسی‌های ترکیبی):")
                p_tab1 = st.checkbox("🛒 ثبت سفارش", value=True)
                p_tab2 = st.checkbox("👤 کارتابل مشتریان", value=True)
                p_tab3 = st.checkbox("📊 مدیریت سفارشات", value=True)
                p_tab4 = st.checkbox("🧠 برنامه‌ریزی تولید (MPS)", value=False)
                p_tab5 = st.checkbox("🏭 تولید و کارگاه", value=False)
                p_tab6 = st.checkbox("📦 انباردار و واحد برش", value=False)
                p_tab7 = st.checkbox("🔍 کنترل کیفیت (QC)", value=False)
                p_tab8 = st.checkbox("📈 گزارشات کلیدی", value=False)
                p_tab9 = st.checkbox("👥 مدیریت کاربران", value=False)
                
                if st.form_submit_button("💾 ذخیره دسترسی‌های کاربر"):
                    if u_username and u_name:
                        df_users_db = get_users_db()
                        
                        permissions_list = []
                        if p_tab1: permissions_list.append("ثبت سفارش")
                        if p_tab2: permissions_list.append("کارتابل مشتریان")
                        if p_tab3: permissions_list.append("مدیریت سفارشات")
                        if p_tab4: permissions_list.append("برنامه‌ریزی تولید")
                        if p_tab5: permissions_list.append("تولید و کارگاه")
                        if p_tab6: permissions_list.append("انباردار")
                        if p_tab7: permissions_list.append("کنترل کیفیت")
                        if p_tab8: permissions_list.append("گزارشات کلیدی")
                        if p_tab9: permissions_list.append("مدیریت کاربران")
                        
                        perms_str = ",".join(permissions_list)

                        if u_username in df_users_db['username'].astype(str).values:
                            df_users_db.loc[df_users_db['username'].astype(str) == u_username, 'password'] = str(u_password)
                            df_users_db.loc[df_users_db['username'].astype(str) == u_username, 'permissions'] = perms_str
                            df_users_db.loc[df_users_db['username'].astype(str) == u_username, 'name'] = str(u_name)
                        else:
                            new_u_df = pd.DataFrame([{
                                "username": u_username, 
                                "password": u_password, 
                                "role": "ترکیبی", 
                                "name": u_name, 
                                "permissions": perms_str
                            }])
                            df_users_db = pd.concat([df_users_db, new_u_df], ignore_index=True)
                        
                        df_users_db.to_excel(USERS_DB_FILE, index=False)
                        log_activity("مدیریت کاربران", f"دسترسی‌های کاربر {u_username} به‌روزرسانی شد.")
                        st.success(f"دسترسی‌های کاربر [{u_name}] با موفقیت ذخیره شد! ✅")
                        st.rerun()
                    else:
                        st.error("لطفاً نام کاربری و نام را وارد کنید.")

            st.markdown("---")
            st.markdown("##### 📋 لیست کاربران و دسترسی‌های فعلی:")
            df_curr_users = get_users_db()
            st.dataframe(df_curr_users[['username', 'name', 'permissions']], use_container_width=True, hide_index=True)